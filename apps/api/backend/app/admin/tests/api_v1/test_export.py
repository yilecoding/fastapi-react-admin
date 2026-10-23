"""导出（#11 · v0.4）。

盯的是**「文件生成了、内容是错的」**那一类 —— 它们都不报错：

- 导出的行和列表看到的对不上（筛选参数静默丢了、数据范围在 worker 里没生效）
- 19 位 ID 在文件里塌成科学计数
- 攻击者填进日志的 `=…` 在 Excel 里成了公式
- 超过上限被悄悄截断
- 别人的导出能被下载

后台导出（Celery）这条路不起 worker：把投递换成记账，再在 TestClient 自己的事件循环里
（`client.portal`）直接调 `export_service.execute` —— 和 worker 跑的是同一个函数，
而且那个事件循环里**没有请求上下文**，正好是要验的「worker 里数据范围靠什么生效」。
"""

import asyncio
import csv
import io
import uuid
import zipfile

from collections.abc import Iterator
from datetime import timedelta
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import bcrypt
import pytest

from openpyxl import load_workbook
from sqlalchemy import delete, insert, text, update
from starlette.testclient import TestClient

from backend.app.admin.model import DataRule, DataScope, Dept, ExportJob, LoginLog, Role, User
from backend.app.admin.model.m2m import data_scope_rule, role_data_scope, role_menu, user_role
from backend.app.admin.service.export_service import export_service
from backend.app.admin.utils.password_security import get_hash_password
from backend.common.enums import ExportJobStatus
from backend.common.enums import RoleDataRuleExpressionType as Expr
from backend.common.enums import RoleDataRuleOperatorType as Op
from backend.core.conf import settings
from backend.database.db import create_database_async_engine, get_database_url
from backend.utils.snowflake import snowflake
from backend.utils.timezone import timezone

PASSWORD = 'Exp!123456'
RUN = uuid.uuid4().hex[:6].upper()

#: 种子里的菜单（三个方言同一套 ID）
SYS_USER_MENU_ID = 2049629108245233672
QUERY_SYS_USER_MENU_ID = 2049629108253622330
EXPORT_SYS_USER_MENU_ID = 2049629108253622350

USERS = '/sys/users'
USERS_EXPORT = '/sys/users/export'
LOGIN_LOG_EXPORT = '/logs/login/export'
EXPORTS = '/sys/exports'


# ─────────────────────────────── 夹具 ───────────────────────────────


class Graph:
    """本模块建出来的部门 / 角色 / 用户，测试里按名字取"""

    def __init__(self) -> None:
        self.dept: dict[str, int] = {}
        self.user: dict[str, int] = {}
        self.role: dict[str, int] = {}
        self.rows: dict[Any, list[int]] = {}

    @staticmethod
    def username(key: str) -> str:
        return f'exp{RUN}_{key}'.lower()


async def _build(graph: Graph) -> None:
    """一个只看得见「本部门用户」的导出人 `r`，一个没有导出权限的 `noexp`。

    ```
    X ── X1          r、peer、noexp 在 X；other 在 X1
    规则：User.dept_id == X   →  看得见 X 里的 3 个人，看不见 other 和种子里的人
    ```
    """
    engine = create_database_async_engine(get_database_url(unittest=True))
    now = timezone.now()
    salt = bcrypt.gensalt()
    pwd = get_hash_password(PASSWORD, salt)
    sid = snowflake.generate

    graph.dept = {'X': sid(), 'X1': sid()}
    depts = [
        {
            'id': graph.dept[key],
            'code': f'EXP{RUN}{key}',
            'name': f'导出测试{key}',
            'sort': 0,
            'leader_id': None,
            'phone': None,
            'email': None,
            'status': 1,
            'parent_id': graph.dept['X'] if key == 'X1' else None,
            'deleted': 0,
            'deleted_time': None,
            'created_time': now,
            'updated_time': None,
        }
        for key in ('X', 'X1')
    ]
    rule_id, scope_id = sid(), sid()
    rules = [
        {
            'id': rule_id,
            'name': f'EXP{RUN}-rule',
            'model': 'User',
            'column': 'dept_id',
            'operator': int(Op.AND),
            'expression': int(Expr.eq),
            'value': str(graph.dept['X']),
            'deleted': 0,
            'deleted_time': None,
            'created_time': now,
            'updated_time': None,
        }
    ]
    scopes = [
        {
            'id': scope_id,
            'name': f'EXP{RUN}-scope',
            'status': 1,
            'deleted': 0,
            'deleted_time': None,
            'created_time': now,
            'updated_time': None,
        }
    ]
    graph.role = {'exporter': sid(), 'viewer': sid()}
    roles = [
        {
            'id': graph.role[key],
            'code': f'EXP{RUN}{key[:3].upper()}',
            'name': f'EXP{RUN}-{key}',
            'status': 1,
            'is_filter_scopes': True,
            'remark': None,
            'deleted': 0,
            'deleted_time': None,
            'created_time': now,
            'updated_time': None,
        }
        for key in ('exporter', 'viewer')
    ]
    r2m = [
        {'id': sid(), 'role_id': graph.role['exporter'], 'menu_id': menu}
        for menu in (SYS_USER_MENU_ID, QUERY_SYS_USER_MENU_ID, EXPORT_SYS_USER_MENU_ID)
    ] + [
        {'id': sid(), 'role_id': graph.role['viewer'], 'menu_id': menu}
        for menu in (SYS_USER_MENU_ID, QUERY_SYS_USER_MENU_ID)
    ]
    r2s = [{'id': sid(), 'role_id': role, 'data_scope_id': scope_id} for role in graph.role.values()]
    s2r = [{'id': sid(), 'data_scope_id': scope_id, 'data_rule_id': rule_id}]

    users, u2r = [], []
    for key, dept, role in (
        ('r', 'X', 'exporter'),
        ('peer', 'X', None),
        ('other', 'X1', None),
        ('noexp', 'X', 'viewer'),
    ):
        pk = sid()
        graph.user[key] = pk
        users.append({
            'id': pk,
            'uuid': str(uuid.uuid4()),
            'username': graph.username(key),
            'nickname': f'导出{key}',
            'password': pwd,
            'salt': salt,
            'email': None,
            'phone': None,
            'avatar': None,
            'timezone': 'Asia/Shanghai',
            'status': 1,
            'is_superuser': False,
            'is_staff': True,
            'join_time': now,
            'last_login_time': now,
            'last_password_changed_time': now,
            'dept_id': graph.dept[dept],
            'deleted': 0,
            'deleted_time': None,
            'created_time': now,
            'updated_time': None,
        })
        if role:
            u2r.append({'id': sid(), 'user_id': pk, 'role_id': graph.role[role]})

    async with engine.begin() as conn:
        await conn.execute(insert(Dept.__table__), depts)
        await conn.execute(insert(DataRule.__table__), rules)
        await conn.execute(insert(DataScope.__table__), scopes)
        await conn.execute(insert(Role.__table__), roles)
        await conn.execute(insert(User.__table__), users)
        await conn.execute(insert(data_scope_rule), s2r)
        await conn.execute(insert(role_data_scope), r2s)
        await conn.execute(insert(user_role), u2r)
        await conn.execute(insert(role_menu), r2m)
    await engine.dispose()

    graph.rows = {
        user_role: [r['id'] for r in u2r],
        role_menu: [r['id'] for r in r2m],
        role_data_scope: [r['id'] for r in r2s],
        data_scope_rule: [r['id'] for r in s2r],
        User.__table__: list(graph.user.values()),
        Role.__table__: list(graph.role.values()),
        DataScope.__table__: [scope_id],
        DataRule.__table__: [rule_id],
        Dept.__table__: list(graph.dept.values()),
    }


async def _sql(statement: str, params: dict[str, Any] | None = None) -> None:
    """独立引擎上执行一句（不能复用绑在 TestClient 循环上的会话）"""
    engine = create_database_async_engine(get_database_url(unittest=True))
    try:
        async with engine.begin() as conn:
            await conn.execute(text(statement), params or {})
    finally:
        await engine.dispose()


async def _update_job(job_id: str, **values: Any) -> None:
    """改一条导出记录。

    🔴 **带时间的值走 Core（列类型 `TimeZone` 管绑定），不能写在 `text()` 里。** SQL Server（aioodbc）会把
    带时区的 datetime 的偏移**丢掉**、按 `+00:00` 存：「一分钟前」存成了八小时之后，过期 / 卡住的判定全都不成立 ——
    PostgreSQL 上绿、SQL Server 上红（实测：`20:16 +08:00` 用 `text()` 写进去读回来是 `20:16 +00:00`）
    """
    engine = create_database_async_engine(get_database_url(unittest=True))
    try:
        async with engine.begin() as conn:
            table = ExportJob.__table__
            await conn.execute(update(table).where(table.c.id == int(job_id)).values(**values))
    finally:
        await engine.dispose()


async def _teardown(graph: Graph) -> None:
    engine = create_database_async_engine(get_database_url(unittest=True))
    async with engine.begin() as conn:
        for table, ids in graph.rows.items():
            if ids:
                await conn.execute(delete(table).where(table.c.id.in_(ids)))
    await engine.dispose()


@pytest.fixture(scope='module')
def g(client: TestClient) -> Iterator[Graph]:
    graph = Graph()
    asyncio.run(_build(graph))
    try:
        yield graph
    finally:
        asyncio.run(_teardown(graph))


@pytest.fixture(autouse=True)
def isolated_upload_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """导出文件落在内置本地存储的私有树上，顶到 tmp（接缝同 `test_file.py`：`file_ops.upload_root()`）"""
    from backend.utils import file_ops

    upload_dir = tmp_path / 'upload'
    public_dir = tmp_path / 'upload-public'
    upload_dir.mkdir()
    public_dir.mkdir()
    monkeypatch.setattr(file_ops, 'UPLOAD_DIR', upload_dir)
    monkeypatch.setattr(file_ops, 'PUBLIC_UPLOAD_DIR', public_dir)
    return upload_dir


@pytest.fixture(autouse=True)
def _purge_jobs() -> Iterator[None]:
    """每条测试收尾硬删导出记录：并发上限按「排队 / 生成中」的条数算，残留会让下一条测试撞上限"""
    yield
    asyncio.run(_sql('DELETE FROM sys_export_job'))


@pytest.fixture
def queue(monkeypatch: pytest.MonkeyPatch) -> list[list[Any]]:
    """把 Celery 投递换成记账，并强制走后台路径"""
    from backend.app.task.celery import celery_app

    sent: list[list[Any]] = []

    def fake_send_task(name: str, args: list[Any] | None = None, **_: Any) -> SimpleNamespace:
        assert name == 'export.run'
        sent.append(args or [])
        return SimpleNamespace(id=f'fake-{len(sent)}')

    monkeypatch.setattr(celery_app, 'send_task', fake_send_task)
    monkeypatch.setattr(settings, 'EXPORT_SYNC_MAX_ROWS', 0)
    return sent


def login(client: TestClient, username: str) -> dict[str, str]:
    resp = client.post('/auth/login/swagger', params={'username': username, 'password': PASSWORD})
    assert resp.status_code == 200, resp.text
    body = resp.json()
    return {'Authorization': f'{body["token_type"]} {body["access_token"]}'}


def start(client: TestClient, headers: dict[str, str], path: str, body: dict[str, Any] | None = None) -> dict[str, Any]:
    """发起导出，返回信封（被拒时 HTTP 是 400、`code` 也是 400，调用方看 `code`）"""
    resp = client.post(path, headers=headers, json=body or {})
    body_json = resp.json()
    assert resp.status_code == body_json['code'] or body_json['code'] == 200, resp.text
    return body_json


def download(client: TestClient, headers: dict[str, str], job_id: str) -> bytes:
    resp = client.get(f'{EXPORTS}/{job_id}/download', headers=headers)
    assert resp.status_code == 200, resp.text
    return resp.content


def sheet_rows(raw: bytes) -> tuple[list[str], list[list[Any]]]:
    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    rows = [list(row) for row in ws.iter_rows(values_only=True)]
    return [str(cell) for cell in rows[0]], rows[1:]


def sheet_xml(raw: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(raw)) as zf:
        return zf.read('xl/worksheets/sheet1.xml').decode()


def run_worker(client: TestClient, job_id: str) -> None:
    """在 TestClient 的事件循环里跑 worker 的那个函数（共享的会话 / Redis 都绑在这个循环上）"""
    client.portal.call(export_service.execute, int(job_id))


def all_user_ids(client: TestClient, headers: dict[str, str], **params: Any) -> set[str]:
    resp = client.get(USERS, headers=headers, params={'size': 200, **params})
    body = resp.json()
    assert body['code'] == 200, body
    assert body['data']['total'] <= 200, '用户超过一页，这个辅助函数要改成翻页'
    return {item['id'] for item in body['data']['items']}


# ─────────────────────────────── 小批量：当场生成 ───────────────────────────────


def test_sync_export_matches_the_list(client: TestClient, token_headers: dict[str, str]) -> None:
    expected = all_user_ids(client, token_headers)
    assert len(expected) >= 2  # 先断言「有」

    job = start(client, token_headers, USERS_EXPORT)['data']
    assert job['status'] == ExportJobStatus.succeeded
    assert job['total'] == job['row_count'] == len(expected)
    assert job['file_name'].endswith('.xlsx')
    assert len(job['file_sha256']) == 64

    header, rows = sheet_rows(download(client, token_headers, job['id']))
    assert header[0] == '用户名'
    assert header[-1] == '用户 ID'
    assert {row[-1] for row in rows} == expected


def test_snowflake_ids_survive_as_text(client: TestClient, token_headers: dict[str, str]) -> None:
    job = start(client, token_headers, USERS_EXPORT, {'columns': ['id', 'username']})['data']
    raw = download(client, token_headers, job['id'])
    _, rows = sheet_rows(raw)
    ids = [row[0] for row in rows]
    assert ids and all(isinstance(pk, str) and len(pk) == 19 for pk in ids), ids
    # 文件里就是文本，没有一个被写成数值（`t="n"` 写进去的那一刻就丢了精度）
    assert '<c r="A2" t="n"' not in sheet_xml(raw)


def test_filters_are_the_list_filters(client: TestClient, token_headers: dict[str, str]) -> None:
    expected = all_user_ids(client, token_headers, username='admin')
    everyone = all_user_ids(client, token_headers)
    assert expected and len(expected) < len(everyone)  # 筛选真的筛掉了人，否则下面的断言分不出对错

    job = start(client, token_headers, USERS_EXPORT, {'filters': {'username': 'admin'}})['data']
    _, rows = sheet_rows(download(client, token_headers, job['id']))
    assert {row[-1] for row in rows} == expected


def test_an_unknown_filter_is_rejected_not_ignored(client: TestClient, token_headers: dict[str, str]) -> None:
    """拼错的参数被忽略 = 导出的是全表，而用户以为是筛过的那批"""
    body = start(client, token_headers, USERS_EXPORT, {'filters': {'user_name': 'admin'}})
    assert body['code'] == 400
    assert 'user_name' in body['msg']


def test_columns_follow_the_requested_order(client: TestClient, token_headers: dict[str, str]) -> None:
    job = start(client, token_headers, USERS_EXPORT, {'columns': ['email', 'username', 'no_such_column']})['data']
    header, _ = sheet_rows(download(client, token_headers, job['id']))
    assert header == ['邮箱', '用户名']


def test_selected_ids_only(client: TestClient, token_headers: dict[str, str]) -> None:
    one = min(all_user_ids(client, token_headers))
    job = start(client, token_headers, USERS_EXPORT, {'ids': [one]})['data']
    assert job['total'] == 1
    _, rows = sheet_rows(download(client, token_headers, job['id']))
    assert [row[-1] for row in rows] == [one]


def test_empty_result_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    body = start(client, token_headers, USERS_EXPORT, {'filters': {'username': f'no-such-{RUN}'}})
    assert body['code'] == 400


def test_over_the_limit_is_rejected_not_truncated(
    client: TestClient, token_headers: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    total = len(all_user_ids(client, token_headers))
    assert total >= 2
    monkeypatch.setattr(settings, 'EXPORT_MAX_ROWS', total - 1)
    body = start(client, token_headers, USERS_EXPORT)
    assert body['code'] == 400
    assert str(total) in body['msg']
    # 拒绝就是拒绝：不留一条「生成了 N-1 行」的记录
    listed = client.get(EXPORTS, headers=token_headers).json()['data']
    assert listed['total'] == 0


def test_csv_is_rfc4180_with_bom(client: TestClient, token_headers: dict[str, str]) -> None:
    job = start(client, token_headers, USERS_EXPORT, {'format': 'csv', 'columns': ['username', 'id']})['data']
    assert job['file_name'].endswith('.csv')
    raw = download(client, token_headers, job['id'])
    assert raw.startswith(b'\xef\xbb\xbf')  # 没有 BOM，Excel 按 GBK 打开中文全是乱码
    assert b'\r\n' in raw
    rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
    assert rows[0] == ['用户名', '用户 ID']
    assert all(len(row[1]) == 19 for row in rows[1:])


# ─────────────────────────────── 公式注入 ───────────────────────────────


@pytest.fixture
def poisoned_logins() -> Iterator[list[str]]:
    """失败登录记的是**请求里填的**用户名 —— 没登录的人能往登录日志里写任何东西。

    直接写库（走登录接口要过验证码和限流，而这里要验的是导出那一侧）
    """
    names = [f'=HYPERLINK("https://example.com/?{RUN}","x")', f'  @SUM(1+1){RUN}', f'＝1+1{RUN}']
    now = timezone.now()
    rows = [
        {
            'id': snowflake.generate(),
            'user_uuid': str(uuid.uuid4()),
            'username': name,
            'status': 0,
            'ip': '127.0.0.1',
            'country': None,
            'region': None,
            'city': None,
            'user_agent': None,
            'os': None,
            'browser': None,
            'device': None,
            'msg': 'pytest',
            'login_time': now,
            'created_time': now,
        }
        for name in names
    ]

    async def _insert() -> None:
        engine = create_database_async_engine(get_database_url(unittest=True))
        try:
            async with engine.begin() as conn:
                await conn.execute(insert(LoginLog.__table__), rows)
        finally:
            await engine.dispose()

    asyncio.run(_insert())
    try:
        yield names
    finally:
        asyncio.run(_sql('DELETE FROM sys_login_log WHERE msg = :m', {'m': 'pytest'}))


def test_formula_payloads_stay_text_in_xlsx(
    client: TestClient, token_headers: dict[str, str], poisoned_logins: list[str]
) -> None:
    job = start(client, token_headers, LOGIN_LOG_EXPORT, {'filters': {'username': RUN}})['data']
    assert job['total'] == len(poisoned_logins)
    raw = download(client, token_headers, job['id'])

    xml = sheet_xml(raw)
    assert '<f>' not in xml  # 文件里没有任何公式元素，打开时不会求值
    _, rows = sheet_rows(raw)
    # 值原样保留（xlsx 不改值），靠单元格类型 + quotePrefix 挡
    assert {row[1] for row in rows} == set(poisoned_logins)
    ws = load_workbook(io.BytesIO(raw)).worksheets[0]
    assert all(ws.cell(row=r, column=2).quotePrefix for r in range(2, len(rows) + 2))


def test_formula_payloads_are_quoted_in_csv(
    client: TestClient, token_headers: dict[str, str], poisoned_logins: list[str]
) -> None:
    job = start(client, token_headers, LOGIN_LOG_EXPORT, {'format': 'csv', 'filters': {'username': RUN}})['data']
    rows = list(csv.reader(io.StringIO(download(client, token_headers, job['id']).decode('utf-8-sig'))))
    names = {row[1] for row in rows[1:]}
    # 前导空白后面的 `@`、全角的 `＝` 也算（EspoCRM CVE-2026-55154 就是垫空白绕过的）
    assert names == {f"'{name}" for name in poisoned_logins}


# ─────────────────────────────── 大批量：后台 ───────────────────────────────


def test_background_export_is_scoped_to_the_requester(
    client: TestClient, g: Graph, queue: list[list[Any]], token_headers: dict[str, str]
) -> None:
    """🔴 worker 里没有请求，数据范围的上下文是空的 —— 空 = 不过滤。

    导出人只看得见本部门的 3 个人；不锚定执行身份的话，worker 会导出全库
    """
    headers = login(client, g.username('r'))
    visible = all_user_ids(client, headers)
    assert visible == {str(g.user[key]) for key in ('r', 'peer', 'noexp')}
    assert len(all_user_ids(client, token_headers)) > len(visible)  # 管理员看得更多，才分得出对错

    job = start(client, headers, USERS_EXPORT)['data']
    assert job['status'] == ExportJobStatus.pending
    assert job['total'] == 3
    assert queue == [[job['id']]]

    run_worker(client, job['id'])

    done = client.get(f'{EXPORTS}/{job["id"]}', headers=headers).json()['data']
    assert done['status'] == ExportJobStatus.succeeded, done
    _, rows = sheet_rows(download(client, headers, job['id']))
    assert {row[-1] for row in rows} == visible


def test_worker_rechecks_permission_before_running(client: TestClient, g: Graph, queue: list[list[Any]]) -> None:
    """排队期间导出权限被收回，任务不该照跑"""
    headers = login(client, g.username('r'))
    job = start(client, headers, USERS_EXPORT)['data']
    assert job['status'] == ExportJobStatus.pending

    asyncio.run(
        _sql(
            'DELETE FROM sys_role_menu WHERE role_id = :r AND menu_id = :m',
            {'r': g.role['exporter'], 'm': EXPORT_SYS_USER_MENU_ID},
        )
    )
    # 用户快照在 Redis 里，改了角色菜单要清（接口改授权时 service 会做这件事）
    client.portal.call(_clear_user_cache, g.user['r'])
    try:
        run_worker(client, job['id'])
        done = client.get(f'{EXPORTS}/{job["id"]}', headers=headers).json()['data']
        assert done['status'] == ExportJobStatus.failed
        assert '权限' in done['error']
    finally:
        asyncio.run(
            _sql(
                'INSERT INTO sys_role_menu (id, role_id, menu_id) VALUES (:id, :r, :m)',
                {'id': snowflake.generate(), 'r': g.role['exporter'], 'm': EXPORT_SYS_USER_MENU_ID},
            )
        )
        client.portal.call(_clear_user_cache, g.user['r'])


async def _clear_user_cache(user_id: int) -> None:
    from backend.database.redis import redis_client

    await redis_client.delete(f'{settings.JWT_USER_REDIS_PREFIX}:{user_id}')


def test_queue_down_marks_the_job_failed(
    client: TestClient, token_headers: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    """投不出去就当场失败，而不是留一条永远「排队」的记录"""
    from backend.app.task.celery import celery_app

    def broken(*_: Any, **__: Any) -> None:
        raise ConnectionError('broker down')

    monkeypatch.setattr(celery_app, 'send_task', broken)
    monkeypatch.setattr(settings, 'EXPORT_SYNC_MAX_ROWS', 0)
    job = start(client, token_headers, USERS_EXPORT)['data']
    assert job['status'] == ExportJobStatus.failed
    assert job['error']


def test_concurrent_exports_are_capped(
    client: TestClient, token_headers: dict[str, str], queue: list[list[Any]], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, 'EXPORT_MAX_CONCURRENT_PER_USER', 1)
    first = start(client, token_headers, USERS_EXPORT)['data']
    assert first['status'] == ExportJobStatus.pending
    second = start(client, token_headers, USERS_EXPORT)
    assert second['code'] == 400
    assert len(queue) == 1


# ─────────────────────────────── 记录：只给本人 ───────────────────────────────


def test_records_are_private(client: TestClient, token_headers: dict[str, str], g: Graph) -> None:
    job = start(client, token_headers, USERS_EXPORT)['data']
    stranger = login(client, g.username('r'))
    # 一律 404 而不是 403：403 等于告诉对方「这个 ID 存在」
    assert client.get(f'{EXPORTS}/{job["id"]}', headers=stranger).json()['code'] == 404
    assert client.get(f'{EXPORTS}/{job["id"]}/download', headers=stranger).status_code == 404
    assert client.delete(f'{EXPORTS}/{job["id"]}', headers=stranger).json()['code'] == 404
    assert client.get(EXPORTS, headers=stranger).json()['data']['total'] == 0
    # 本人还拿得到（前面的拒绝不是因为记录坏了）
    assert client.get(f'{EXPORTS}/{job["id"]}/download', headers=token_headers).status_code == 200


def test_export_needs_its_own_permission(client: TestClient, g: Graph) -> None:
    """能看列表 ≠ 能导出：导出是批量拿走数据，权限码单独一个"""
    headers = login(client, g.username('noexp'))
    assert len(all_user_ids(client, headers)) == 3  # 看得见列表（同部门的 r、peer、noexp 自己）
    resp = client.post(USERS_EXPORT, headers=headers, json={})
    assert resp.status_code == 403


def test_expired_file_is_refused_then_purged(
    client: TestClient, token_headers: dict[str, str], isolated_upload_dir: Path
) -> None:
    job = start(client, token_headers, USERS_EXPORT)['data']
    files = list(isolated_upload_dir.rglob('*.xlsx'))
    assert len(files) == 1  # 先断言文件真的落盘了

    asyncio.run(_update_job(job['id'], expire_time=timezone.now() - timedelta(minutes=1)))
    assert client.get(f'{EXPORTS}/{job["id"]}/download', headers=token_headers).json()['code'] == 400

    client.portal.call(_purge)
    done = client.get(f'{EXPORTS}/{job["id"]}', headers=token_headers).json()['data']
    assert done['status'] == ExportJobStatus.expired
    assert not files[0].exists()


async def _purge() -> None:
    # ⚠️ 用测试库的会话。直接 `from backend.database.db import async_db_session` 拿到的是**开发库**
    # （conftest 只替换了被测模块里的那个名字）—— 第一版就这么写错过，撞上开发库里那张旧表才发现
    from backend.tests.utils.db import async_test_db_session

    async with async_test_db_session() as db:
        await export_service.purge(db=db)


def test_a_job_stuck_running_is_settled_as_failed(client: TestClient, token_headers: dict[str, str]) -> None:
    """worker 死在半路：记录不能永远停在「生成中」"""
    job = start(client, token_headers, USERS_EXPORT)['data']
    long_ago = timezone.now() - timedelta(seconds=settings.CELERY_TASK_TIME_LIMIT + 600)
    asyncio.run(_update_job(job['id'], status=ExportJobStatus.running.value, started_time=long_ago))
    items = client.get(EXPORTS, headers=token_headers).json()['data']['items']
    mine = next(item for item in items if item['id'] == job['id'])
    assert mine['status'] == ExportJobStatus.failed
    assert mine['error']


def test_info_sheet_says_who_exported(client: TestClient, token_headers: dict[str, str]) -> None:
    job = start(client, token_headers, USERS_EXPORT, {'filters': {'username': 'admin'}})['data']
    wb = load_workbook(io.BytesIO(download(client, token_headers, job['id'])))
    assert len(wb.worksheets) == 2
    info = {row[0]: row[1] for row in wb.worksheets[1].iter_rows(values_only=True)}
    assert 'admin' in info['导出人']
    assert 'admin' in info['筛选条件']
    assert wb.properties.creator == 'admin'


# ─────────────────────────────── 每个资源 ───────────────────────────────


@pytest.fixture
def some_logs() -> Iterator[None]:
    """刚重建的测试库里日志是空的（swagger 登录不记登录日志、操作日志是后台消费者落库的），垫几行"""
    from backend.app.admin.model import OperaLog

    now = timezone.now()
    logins = [
        {
            'id': snowflake.generate(),
            'user_uuid': str(uuid.uuid4()),
            'username': f'pytest-log-{RUN}',
            'status': 1,
            'ip': '127.0.0.1',
            'country': None,
            'region': None,
            'city': None,
            'user_agent': None,
            'os': None,
            'browser': None,
            'device': None,
            'msg': 'pytest',
            'login_time': now,
            'created_time': now,
        }
        for _ in range(2)
    ]
    operas = [
        {
            'id': snowflake.generate(),
            'trace_id': uuid.uuid4().hex,
            'username': f'pytest-log-{RUN}',
            'method': 'GET',
            'title': 'pytest',
            'path': '/api/v1/pytest',
            'ip': '127.0.0.1',
            'country': None,
            'region': None,
            'city': None,
            'user_agent': None,
            'os': None,
            'browser': None,
            'device': None,
            'args': None,
            'request_headers': None,
            'response_headers': None,
            'response_body': None,
            'status': 1,
            'code': '200',
            'msg': 'pytest',
            'cost_time': 1.5,
            'opera_time': now,
            'created_time': now,
        }
        for _ in range(2)
    ]

    async def _insert() -> None:
        engine = create_database_async_engine(get_database_url(unittest=True))
        try:
            async with engine.begin() as conn:
                await conn.execute(insert(LoginLog.__table__), logins)
                await conn.execute(insert(OperaLog.__table__), operas)
        finally:
            await engine.dispose()

    asyncio.run(_insert())
    try:
        yield
    finally:
        asyncio.run(_sql('DELETE FROM sys_login_log WHERE msg = :m', {'m': 'pytest'}))
        asyncio.run(_sql('DELETE FROM sys_opera_log WHERE msg = :m', {'m': 'pytest'}))


@pytest.mark.parametrize(
    ('list_path', 'export_path', 'by_fixture'),
    [
        ('/sys/users', '/sys/users/export', False),
        ('/sys/dict-datas', '/sys/dict-datas/export', False),
        ('/sys/notices', '/sys/notices/export', False),
        # 日志按 `some_logs` 垫的那几行筛：不筛的话，表里有多少行取决于之前跑过什么 ——
        # 跑过一轮 E2E 之后 `fba_test` 的操作日志就有七千多行，超过当场生成的上限，导出转去后台，
        # 这条测试红在「状态是排队中」上，和它要验的东西毫无关系（实测）。而且测试自己的每个请求
        # 都在往操作日志里写，不筛就只能比「至少」
        ('/logs/login', '/logs/login/export', True),
        ('/logs/opera', '/logs/opera/export', True),
    ],
)
def test_every_resource_exports_what_its_list_shows(
    client: TestClient,
    token_headers: dict[str, str],
    list_path: str,
    export_path: str,
    by_fixture: bool,  # ruff:ignore[boolean-type-hint-positional-argument]
    some_logs: None,
) -> None:
    filters = {'username': f'pytest-log-{RUN}'} if by_fixture else {}
    listed = client.get(list_path, headers=token_headers, params={'size': 1, **filters}).json()
    assert listed['code'] == 200, listed
    total = listed['data']['total']
    assert total > 0  # 先断言「有」：空表导出被拒，分不出「导出坏了」还是「没数据」

    job = start(client, token_headers, export_path, {'filters': filters})['data']
    assert job['status'] == ExportJobStatus.succeeded, job
    _, rows = sheet_rows(download(client, token_headers, job['id']))
    # 快照语义：数出来多少行就导多少行
    assert len(rows) == job['total'] == job['row_count'] == total


def test_dept_export_is_in_tree_order(client: TestClient, token_headers: dict[str, str], g: Graph) -> None:
    """上级在前、下级在后，下级那一行带着上级的编码 —— 导出的表改完能原样导回去"""
    job = start(client, token_headers, '/sys/depts/export', {'columns': ['code', 'parent_code', 'id']})['data']
    _, rows = sheet_rows(download(client, token_headers, job['id']))
    position = {row[0]: index for index, row in enumerate(rows)}
    child = f'EXP{RUN}X1'
    assert child in position and f'EXP{RUN}X' in position  # 先断言这棵小树在里面
    assert rows[position[child]][1] == f'EXP{RUN}X'
    for code, parent_code, _ in rows:
        if parent_code:
            assert position[parent_code] < position[code], f'{code} 排在上级 {parent_code} 前面'


def test_every_export_label_is_translated_in_both_languages() -> None:
    """表头 / 取值 / 资源名都是**运行时拼出来的键**（`t(column.label)`），`test_i18n.py` 的正则扫不到。

    缺键的表现：那门语言的人拿到的文件表头是一串 `export.column.xxx`
    """
    import yaml

    from backend.common.dataio.exporting import export_specs, load_dataio_modules
    from backend.core.path_conf import LOCALE_DIR

    load_dataio_modules()
    specs = export_specs()
    assert len(specs) >= 6  # 插件的定义也得登记进来（worker 走的就是这条加载路径）

    keys: set[str] = set()
    for spec in specs.values():
        keys.add(spec.label)
        keys.update(spec.filter_labels.values())
        for column in spec.columns:
            keys.add(column.label)
            keys.update((column.choices or {}).values())

    for name in ('zh-CN.yml', 'en-US.yml'):
        locale = yaml.safe_load((LOCALE_DIR / name).read_text(encoding='utf-8'))
        missing = []
        for key in sorted(keys):
            node: Any = locale
            for part in key.split('.'):
                node = node.get(part) if isinstance(node, dict) else None
            if not isinstance(node, str):
                missing.append(key)
        assert not missing, f'{name} 缺这些导出文案：{missing}'
