"""用户批量导入（预览 → 提交）。

这批用例盯的是**「导入成功了、人数对不上」**那一类问题——它们全都不报错：
文件内部重名、部门编码查不到、token 被重放、排除项没生效。

⚠️ 收尾一律**硬删**建出来的用户。`DELETE /sys/users/{pk}` 是逻辑删除
（`deleted` 置成行自己的 id），行永久留在表里、同名还能再建、测试照样绿 ——
代价是每跑一次 pytest 就往 `fba_test` 里堆一批（同 `conftest.temp_user` 那条）。
"""

import asyncio
import io

import pytest

from openpyxl import Workbook, load_workbook
from starlette.testclient import TestClient

from backend.core.conf import settings
from backend.utils.excel_ops import parse_table

#: 这批测试建出来的用户名前缀，收尾按它清
PREFIX = 'pytest_imp_'

#: 导入用的默认密码。必须过 `validate_password_strength()`（要含字母和数字），
#: 所以**不能**写 `123456` —— 那是种子密码，`is_has_letter` 那条就挡住了
DEFAULT_PASSWORD = 'Pytest@Imp1'


@pytest.fixture(autouse=True)
def _default_password() -> None:
    """给这批用例配上导入默认密码。

    ⚠️ 不能只在某一条里设 —— `USER_IMPORT_DEFAULT_PASSWORD` 默认是空串
    （= 功能关闭），漏设的表现是所有用例一起报「未配置默认密码」，
    看起来像功能坏了。
    """
    original = settings.USER_IMPORT_DEFAULT_PASSWORD
    settings.USER_IMPORT_DEFAULT_PASSWORD = DEFAULT_PASSWORD
    yield
    settings.USER_IMPORT_DEFAULT_PASSWORD = original


@pytest.fixture
def refs(client: TestClient, token_headers: dict[str, str]) -> dict[str, str]:
    """从**接口**取部门/角色编码，不硬编码。

    ⚠️ 三个方言的种子各有一套雪花 ID，写死会让这批测试只在一种库上能跑
    （同 `conftest.temp_user` 的那条注释）。编码倒是跨库一致，但也从接口读，
    免得种子换了占位码之后这里静默指空。
    """
    dept = client.get('/sys/depts', headers=token_headers).json()['data'][0]
    role = client.get('/sys/roles/all', headers=token_headers).json()['data'][0]
    return {'dept_code': dept['code'], 'role_code': role['code']}


@pytest.fixture(autouse=True)
def _purge_imported() -> None:
    """收尾硬删这批用户 + 它们的关联行"""
    yield
    from sqlalchemy import text

    from backend.database.db import create_database_async_engine, create_database_async_session, get_database_url

    async def _purge() -> None:
        # ⚠️ 不能复用共享的 `async_test_db_session`：它绑在 TestClient 的事件循环上，
        # 这里的 `asyncio.run()` 是另起一个 —— 用它会 `attached to a different loop`
        engine = create_database_async_engine(get_database_url(unittest=True))
        session_maker = create_database_async_session(engine)
        try:
            async with session_maker.begin() as session:
                await session.execute(
                    text('DELETE FROM sys_user_role WHERE user_id IN (SELECT id FROM sys_user WHERE username LIKE :p)'),
                    {'p': f'{PREFIX}%'},
                )
                await session.execute(text('DELETE FROM sys_user WHERE username LIKE :p'), {'p': f'{PREFIX}%'})
        finally:
            await engine.dispose()

    asyncio.run(_purge())


def build_xlsx(rows: list[list], header: list[str] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header or ['username', 'nickname', 'email', 'phone', 'dept_code', 'role_codes'])
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client: TestClient, headers: dict[str, str], raw: bytes, name: str = 'u.xlsx') -> dict:
    res = client.post(
        '/sys/users/import/preview',
        headers=headers,
        files={'file': (name, raw, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
    )
    return res.json()


# --------------------------------------------------------------------------- 模板


def test_template_downloads_and_parses_back_through_our_own_parser(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    """模板和解析器分叉是无声的：下发一份自己都读不回来的模板，没有任何一处会报错"""
    res = client.get('/sys/users/import/template', headers=token_headers)
    assert res.status_code == 200
    assert 'attachment' in res.headers['content-disposition']

    wb = load_workbook(io.BytesIO(res.content))
    assert wb.worksheets[0].title == 'data'
    # 说明在第二页，而**活动页签必须是数据页** —— 否则用户随手另存之后
    # `wb.active` 就指向说明页了（见 excel_ops 里 `_pick_sheet` 那条）
    assert wb.active.title == 'data'
    assert 'README' in wb.sheetnames


# --------------------------------------------------------------------------- 权限


def test_a_non_superuser_is_rejected_by_all_three_endpoints(client: TestClient) -> None:
    """🔴 **前端藏了按钮不等于接口有校验**，这两件事是分开写的，漏挂依赖不会在
    任何地方报错。批量建号的影响面比单条建号更大，三条都必须挡住。

    用种子里的 `test`（STAFF 角色、非超管），同
    `security/test_read_endpoint_rbac_gaps.py` 的做法。
    """
    resp = client.post('/auth/login/swagger', params={'username': 'test', 'password': '123456'})
    resp.raise_for_status()
    body = resp.json()
    headers = {'Authorization': f'{body["token_type"]} {body["access_token"]}'}

    assert client.get('/sys/users/import/template', headers=headers).status_code == 403
    assert (
        client.post(
            '/sys/users/import/preview',
            headers=headers,
            files={'file': ('u.xlsx', build_xlsx([]), 'application/octet-stream')},
        ).status_code
        == 403
    )
    assert client.post('/sys/users/import/commit', headers=headers, json={'import_token': 'x'}).status_code == 403


# --------------------------------------------------------------------------- 预览


def test_preview_accepts_a_clean_file(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    raw = build_xlsx([[f'{PREFIX}a', '甲', '', '13800138001', refs['dept_code'], refs['role_code']]])
    body = upload(client, token_headers, raw)
    assert body['code'] == 200, body
    assert body['data']['total'] == 1
    assert body['data']['valid'] == 1
    assert body['data']['rows'][0]['ok'] is True
    assert body['data']['import_token']


def test_preview_does_not_create_anything(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    """🔴 预览必须是纯读。它落库的话「先看看」这个动作本身就产生了用户"""
    raw = build_xlsx([[f'{PREFIX}dry', '', '', '', refs['dept_code'], refs['role_code']]])
    upload(client, token_headers, raw)
    listed = client.get('/sys/users', headers=token_headers, params={'username': f'{PREFIX}dry'}).json()
    assert listed['data']['total'] == 0


def test_unknown_dept_code_marks_only_that_row(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    """一行错不该拖垮整份文件 —— 这正是两阶段设计要买到的东西"""
    raw = build_xlsx([
        [f'{PREFIX}ok', '', '', '', refs['dept_code'], refs['role_code']],
        [f'{PREFIX}bad', '', '', '', 'NO_SUCH_DEPT', refs['role_code']],
    ])
    data = upload(client, token_headers, raw)['data']
    assert (data['total'], data['valid']) == (2, 1)
    bad = next(row for row in data['rows'] if row['username'] == f'{PREFIX}bad')
    assert not bad['ok'] and any('NO_SUCH_DEPT' in msg for msg in bad['errors'])


def test_unknown_role_code_marks_only_that_row(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    raw = build_xlsx([[f'{PREFIX}r', '', '', '', refs['dept_code'], 'NO_SUCH_ROLE']])
    data = upload(client, token_headers, raw)['data']
    assert data['valid'] == 0
    assert any('NO_SUCH_ROLE' in msg for msg in data['rows'][0]['errors'])


def test_duplicate_inside_the_file_is_caught(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    """🔴 库里没有、文件里有两行同名：逐行对着库查**全都通过**，然后在提交
    阶段撞唯一约束 —— 而那时候第一行已经建出来了。所以文件内部这一维必须单独查。
    """
    raw = build_xlsx([
        [f'{PREFIX}dup', '', '', '', refs['dept_code'], refs['role_code']],
        [f'{PREFIX}dup', '', '', '', refs['dept_code'], refs['role_code']],
    ])
    data = upload(client, token_headers, raw)['data']
    assert data['valid'] == 1, '第二行应该被判重复'
    assert data['rows'][1]['ok'] is False


def test_duplicate_against_the_database_is_caught(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """库里已有的用户名 —— `admin` 一定存在，拿它当探针"""
    raw = build_xlsx([['admin', '', '', '', refs['dept_code'], refs['role_code']]])
    data = upload(client, token_headers, raw)['data']
    assert data['valid'] == 0


def test_missing_required_header_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    body = upload(client, token_headers, build_xlsx([['x']], header=['username']))
    assert body['code'] != 200
    assert 'dept_code' in body['msg']


def test_unrecognized_header_is_surfaced(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    """`phone` 拼错时那一列的数据全丢，而整份文件照样解析成功 —— 纯静默"""
    raw = build_xlsx(
        [[f'{PREFIX}h', '', '', '13800138002', refs['dept_code'], refs['role_code']]],
        header=['username', 'nickname', 'email', 'phone_number', 'dept_code', 'role_codes'],
    )
    data = upload(client, token_headers, raw)['data']
    assert data['ignored_headers'] == ['phone_number']
    assert data['rows'][0]['phone'] is None


def test_non_xlsx_upload_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    body = upload(client, token_headers, b'username,dept_code\na,b\n', name='u.csv')
    assert body['code'] != 200 and 'xlsx' in body['msg']


def test_import_is_off_when_no_default_password_is_configured(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """空串 = 功能关闭，不是「用某个内置默认值」"""
    settings.USER_IMPORT_DEFAULT_PASSWORD = ''
    raw = build_xlsx([[f'{PREFIX}off', '', '', '', refs['dept_code'], refs['role_code']]])
    body = upload(client, token_headers, raw)
    assert body['code'] != 200
    # ⚠️ **必须断言到具体那句话。** 只断言「不是 200」时，把这道检查整个去掉
    # 用例照样绿 —— 空密码会接着撞上强度校验，同样报错，只是报的是另一回事
    # （变异检验抓出来的：摘掉这道检查，这条测试原来仍然通过）
    assert 'USER_IMPORT_DEFAULT_PASSWORD' in body['msg'], body['msg']


def test_a_default_password_that_fails_the_policy_is_refused_up_front(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """🔴 默认密码必须自己过强度校验，否则失败是**延迟且分裂**的：导入成功、
    用户能登录，但他去改密码时被「必须含字母」挡住，而两件事看不出是同一个原因。
    `123456` 正是这样一个值（种子密码，无字母）。
    """
    settings.USER_IMPORT_DEFAULT_PASSWORD = '123456'
    raw = build_xlsx([[f'{PREFIX}weak', '', '', '', refs['dept_code'], refs['role_code']]])
    body = upload(client, token_headers, raw)
    assert body['code'] != 200
    # 断言到具体原因，和上面那条区分开 —— 两条都只判「不是 200」的话，
    # 它们会互相顶替，摘掉任何一道检查都还是绿的
    assert '字母' in body['msg'] or 'letter' in body['msg'], body['msg']


# --------------------------------------------------------------------------- 提交


def test_commit_creates_the_users_and_they_can_sign_in(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """端到端那一条：导进去的人**真的能用默认密码登录**。

    只断言「接口返回 200」是不够的 —— 密码 hash 写错、角色没关联、状态是停用，
    这三种都能让接口成功而人登不进来。
    """
    username = f'{PREFIX}login'
    raw = build_xlsx([[username, '登录测试', '', '13800138003', refs['dept_code'], refs['role_code']]])
    token = upload(client, token_headers, raw)['data']['import_token']

    body = client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': token}).json()
    assert body['code'] == 200, body
    assert body['data']['created'] == [username]
    assert body['data']['used_default_password'] is True

    settings.LOGIN_CAPTCHA_ENABLED = False
    login = client.post('/auth/login', json={'username': username, 'password': DEFAULT_PASSWORD})
    assert login.status_code == 200 and login.json()['code'] == 200, login.text


def test_commit_response_never_carries_a_plaintext_password(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """🔴 响应体会被原样写进 `sys_opera_log.response_body`。

    把明文密码放进来 = 永久写进日志表，任何有日志查看权限的人都能翻出来。
    这一条钉住的是「以后有人为了方便把密码加回响应」。
    """
    raw = build_xlsx([[f'{PREFIX}nopwd', '', '', '', refs['dept_code'], refs['role_code']]])
    token = upload(client, token_headers, raw)['data']['import_token']
    res = client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': token})
    assert DEFAULT_PASSWORD not in res.text


def test_excluded_rows_are_not_created(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    raw = build_xlsx([
        [f'{PREFIX}keep', '', '', '', refs['dept_code'], refs['role_code']],
        [f'{PREFIX}drop', '', '', '', refs['dept_code'], refs['role_code']],
    ])
    preview = upload(client, token_headers, raw)['data']
    drop_row = next(row['row_no'] for row in preview['rows'] if row['username'] == f'{PREFIX}drop')

    body = client.post(
        '/sys/users/import/commit',
        headers=token_headers,
        json={'import_token': preview['import_token'], 'exclude_rows': [drop_row]},
    ).json()
    assert body['data']['created'] == [f'{PREFIX}keep']


def test_the_same_token_cannot_be_committed_twice(
    client: TestClient, token_headers: dict[str, str], refs: dict
) -> None:
    """🔴 不作废 token 的话，重放一次就是重复建号 —— 第二次会撞唯一约束报「导入
    失败」，而库里已经多了第一批人。token 在提交前就删掉。
    """
    raw = build_xlsx([[f'{PREFIX}once', '', '', '', refs['dept_code'], refs['role_code']]])
    token = upload(client, token_headers, raw)['data']['import_token']

    first = client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': token}).json()
    assert first['code'] == 200
    second = client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': token}).json()
    assert second['code'] != 200


def test_an_unknown_token_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    body = client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': 'nope'}).json()
    assert body['code'] != 200


def test_roles_are_actually_attached(client: TestClient, token_headers: dict[str, str], refs: dict) -> None:
    """角色没关联上是静默的：用户建出来了、能登录，只是什么都看不见"""
    username = f'{PREFIX}role'
    raw = build_xlsx([[username, '', '', '', refs['dept_code'], refs['role_code']]])
    token = upload(client, token_headers, raw)['data']['import_token']
    client.post('/sys/users/import/commit', headers=token_headers, json={'import_token': token})

    listed = client.get('/sys/users', headers=token_headers, params={'username': username}).json()
    pk = listed['data']['items'][0]['id']
    roles = client.get(f'/sys/users/{pk}/roles', headers=token_headers).json()
    assert [role['code'] for role in roles['data']] == [refs['role_code']]


def test_parsed_template_columns_match_what_the_service_declares() -> None:
    """模板列和解析器声明的列必须是同一份 —— 分叉了只会表现成「某列永远读不到」"""
    from backend.app.admin.service.user_import_service import IMPORT_COLUMNS, user_import_service

    wb = load_workbook(io.BytesIO(user_import_service.template()))
    header = [cell.value for cell in next(wb.worksheets[0].iter_rows(max_row=1))]
    assert header == [column.key for column in IMPORT_COLUMNS]
    # 再走一遍自己的解析器，确认表头能被认全
    filled = io.BytesIO()
    wb.worksheets[0].append(['u', 'n', '', '', 'DEPT_0001', 'ROLE_0001'])
    wb.save(filled)
    assert parse_table(filled.getvalue(), IMPORT_COLUMNS).ignored_headers == []
