"""导入（#11 · v0.4）：部门 · 字典数据 · 框架本身（令牌 / 错误报告 / 全有或全无 / 表头语言）。

用户导入那一份在 `test_user_import.py`（它早于框架，盯的是建号那一侧）。这里盯两类**都不报错**的问题：

- **预览说能导、提交却不对**：预览和提交之间库里变了；同一份文件里先挪走 B 再把 A 挂到 B 下面，
  按库里现在的树查会误报成环、按文件顺序写会在提交时撞上
- **导进去了、树是错的**：上下级成环（整棵子树从部门树里消失）、下级挂到了没建出来的上级下面、
  数据范围外的部门被改掉

⚠️ 收尾一律**硬删**：接口的删除是逻辑删除，编码 / 名称还占着唯一性（同 `test_user_import.py`）。
"""

import asyncio
import io
import re
import uuid

from collections.abc import Iterator
from pathlib import Path
from typing import Any

import bcrypt
import pytest

from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, insert, text
from starlette.testclient import TestClient

from backend.app.admin.model import DataRule, DataScope, Role, User
from backend.app.admin.model.m2m import data_scope_rule, role_data_scope, role_menu, user_role
from backend.app.admin.service.import_service import _table_columns
from backend.app.admin.utils.password_security import get_hash_password
from backend.common.dataio.importing import import_specs
from backend.common.enums import RoleDataRuleExpressionType as Expr
from backend.common.enums import RoleDataRuleOperatorType as Op
from backend.common.i18n import i18n, t
from backend.database.db import create_database_async_engine, get_database_url
from backend.utils.excel_ops import parse_table
from backend.utils.snowflake import snowflake
from backend.utils.timezone import timezone

RUN = uuid.uuid4().hex[:6].upper()
PASSWORD = 'Imp!123456'

DEPTS = '/sys/depts'
DEPT_IMPORT = '/sys/depts/import'
DICT_IMPORT = '/sys/dict-datas/import'
USER_IMPORT = '/sys/users/import'

#: 种子里的菜单（三个方言同一套 ID）：部门管理 · 部门「导入」按钮
SYS_DEPT_MENU_ID = 2049629108245233668
IMPORT_SYS_DEPT_MENU_ID = 2049629108253622352

#: 部门导入模板的表头（中文列名）
DEPT_HEADER = ['部门编码', '部门名称', '上级部门编码', '负责人用户名', '手机号', '邮箱', '排序', '状态']


def code(key: str) -> str:
    """本模块建的部门编码：`IMP<RUN><KEY>`（编码只许大写字母开头、大写字母 / 数字 / 下划线）"""
    return f'IMP{RUN}{key}'


def username(key: str) -> str:
    return f'imp{RUN}_{key}'.lower()


async def _sql(statement: str, params: dict[str, Any] | None = None) -> None:
    """独立引擎上执行一句（不能复用绑在 TestClient 事件循环上的会话）"""
    engine = create_database_async_engine(get_database_url(unittest=True))
    try:
        async with engine.begin() as conn:
            await conn.execute(text(statement), params or {})
    finally:
        await engine.dispose()


@pytest.fixture(autouse=True)
def _purge() -> Iterator[None]:
    """每条收尾硬删本模块建的部门 / 用户 / 字典"""
    yield

    async def _run() -> None:
        engine = create_database_async_engine(get_database_url(unittest=True))
        try:
            async with engine.begin() as conn:
                await conn.execute(text('DELETE FROM sys_dept WHERE code LIKE :p'), {'p': f'IMP{RUN}%'})
                await conn.execute(
                    text('DELETE FROM sys_user_role WHERE user_id IN (SELECT id FROM sys_user WHERE username LIKE :p)'),
                    {'p': f'imp{RUN.lower()}%'},
                )
                await conn.execute(text('DELETE FROM sys_user WHERE username LIKE :p'), {'p': f'imp{RUN.lower()}%'})
                await conn.execute(text('DELETE FROM sys_dict_data WHERE type_code LIKE :p'), {'p': f'imp{RUN}%'})
                await conn.execute(text('DELETE FROM sys_dict_type WHERE code LIKE :p'), {'p': f'imp{RUN}%'})
                await conn.execute(text('DELETE FROM sys_export_job'))
        finally:
            await engine.dispose()

    asyncio.run(_run())


@pytest.fixture(autouse=True)
def isolated_upload_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """往返那一条要先导出：导出文件落在内置本地存储的私有树上，顶到 tmp（同 `test_export.py`）"""
    from backend.utils import file_ops

    upload_dir = tmp_path / 'upload'
    public_dir = tmp_path / 'upload-public'
    upload_dir.mkdir()
    public_dir.mkdir()
    monkeypatch.setattr(file_ops, 'UPLOAD_DIR', upload_dir)
    monkeypatch.setattr(file_ops, 'PUBLIC_UPLOAD_DIR', public_dir)
    return upload_dir


# ─────────────────────────────── 辅助 ───────────────────────────────


def xlsx(header: list[str], rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def dept_rows(*rows: dict[str, Any]) -> bytes:
    """按列 key 写部门行，没给的列留空"""
    keys = ['code', 'name', 'parent_code', 'leader', 'phone', 'email', 'sort', 'status']
    return xlsx(DEPT_HEADER, [[row.get(key) for key in keys] for row in rows])


def preview(client: TestClient, headers: dict[str, str], base: str, raw: bytes, mode: str = 'create') -> dict[str, Any]:
    """预览，返回信封（被拒时 `code` 不是 200，调用方看 `code`）"""
    res = client.post(
        f'{base}/preview',
        headers=headers,
        params={'mode': mode},
        files={'file': ('f.xlsx', raw, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},
    )
    return res.json()


def ok_preview(client: TestClient, headers: dict[str, str], base: str, raw: bytes, mode: str = 'create') -> dict:
    body = preview(client, headers, base, raw, mode)
    assert body['code'] == 200, body
    return body['data']


def commit(client: TestClient, headers: dict[str, str], base: str, token: str, **kwargs: Any) -> dict[str, Any]:
    return client.post(f'{base}/commit', headers=headers, json={'import_token': token, **kwargs}).json()


def row_of(data: dict[str, Any], key: str, value: str) -> dict[str, Any]:
    return next(row for row in data['rows'] if row['values'][key] == value)


def dept_tree(client: TestClient, headers: dict[str, str]) -> dict[str, dict[str, Any]]:
    """本模块建的部门：编码 → 节点（含 `parent_id`），从部门树接口平铺出来"""
    res = client.get(DEPTS, headers=headers, params={'code': f'IMP{RUN}'}).json()
    assert res['code'] == 200, res
    out: dict[str, dict[str, Any]] = {}

    def walk(nodes: list[dict[str, Any]]) -> None:
        for node in nodes:
            out[node['code']] = node
            walk(node.get('children') or [])

    walk(res['data'])
    return out


def make_dept(client: TestClient, headers: dict[str, str], key: str, parent: str | None = None, **extra: Any) -> str:
    """经接口建一个部门，返回 ID"""
    parent_id = dept_tree(client, headers)[code(parent)]['id'] if parent else None
    body = {'code': code(key), 'name': f'导入{RUN}{key}', 'parent_id': parent_id, 'status': 1, **extra}
    res = client.post(DEPTS, headers=headers, json=body).json()
    assert res['code'] == 200, res
    return dept_tree(client, headers)[code(key)]['id']


def make_user(key: str, *, status: int = 1, dept_id: int | None = None) -> int:
    """直接写库建一个用户（只当负责人 / 受限导入人用）"""
    pk = snowflake.generate()
    now = timezone.now()
    salt = bcrypt.gensalt()

    async def _insert() -> None:
        engine = create_database_async_engine(get_database_url(unittest=True))
        try:
            async with engine.begin() as conn:
                await conn.execute(
                    insert(User.__table__),
                    [
                        {
                            'id': pk,
                            'uuid': str(uuid.uuid4()),
                            'username': username(key),
                            'nickname': f'导入{key}',
                            'password': get_hash_password(PASSWORD, salt),
                            'salt': salt,
                            'email': None,
                            'phone': None,
                            'avatar': None,
                            'timezone': 'Asia/Shanghai',
                            'status': status,
                            'is_superuser': False,
                            'is_staff': True,
                            'join_time': now,
                            'last_login_time': now,
                            'last_password_changed_time': now,
                            'dept_id': dept_id,
                            'deleted': 0,
                            'deleted_time': None,
                            'created_time': now,
                            'updated_time': None,
                        }
                    ],
                )
        finally:
            await engine.dispose()

    asyncio.run(_insert())
    return pk


# ═══════════════════════════════ 部门 ═══════════════════════════════


def test_a_tree_in_one_file_is_created_parents_first(client: TestClient, token_headers: dict[str, str]) -> None:
    """下级写在上级前面（人整理表格时不会按层级排）：提交要自己排成先上级后下级"""
    raw = dept_rows(
        {'code': code('C'), 'name': 'C', 'parent_code': code('B')},
        {'code': code('B'), 'name': 'B', 'parent_code': code('A')},
        {'code': code('A'), 'name': f'导入{RUN}A'},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert (data['total'], data['valid'], data['to_create']) == (3, 3, 3), data['rows']

    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    assert body['code'] == 200, body
    assert sorted(body['data']['created']) == sorted([code('A'), code('B'), code('C')])
    assert body['data']['failed'] == []

    tree = dept_tree(client, token_headers)
    assert tree[code('A')]['parent_id'] is None
    assert tree[code('B')]['parent_id'] == tree[code('A')]['id']
    assert tree[code('C')]['parent_id'] == tree[code('B')]['id']


def test_a_cycle_inside_the_file_is_flagged_and_its_subtree_too(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    """X ↔ Y 互为上级：成环的两行都标；挂在 X 下面的 Z 跟着标（它的上级建不出来）；无关的 W 照常"""
    raw = dept_rows(
        {'code': code('X'), 'name': 'X', 'parent_code': code('Y')},
        {'code': code('Y'), 'name': 'Y', 'parent_code': code('X')},
        {'code': code('Z'), 'name': 'Z', 'parent_code': code('X')},
        {'code': code('W'), 'name': f'导入{RUN}W'},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['valid'] == 1
    assert row_of(data, 'code', code('W'))['ok'] is True
    for key in ('X', 'Y'):
        errors = row_of(data, 'code', code(key))['errors']
        assert any('→' in msg and code('X') in msg and code('Y') in msg for msg in errors), errors
    z_errors = row_of(data, 'code', code('Z'))['errors']
    assert any(code('X') in msg for msg in z_errors), z_errors


def test_moving_a_department_under_its_own_child_is_caught_in_preview(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    """🔴 环的一半在库里：P ── Q 已经存在，文件把 P 挂到 Q 下面。只查文件里的行之间是查不出来的"""
    make_dept(client, token_headers, 'P')
    make_dept(client, token_headers, 'Q', parent='P')

    data = ok_preview(
        client, token_headers, DEPT_IMPORT, dept_rows({'code': code('P'), 'parent_code': code('Q')}), 'update'
    )
    row = row_of(data, 'code', code('P'))
    assert row['ok'] is False
    assert any(code('Q') in msg and '→' in msg for msg in row['errors']), row['errors']
    assert dept_tree(client, token_headers)[code('P')]['parent_id'] is None


def test_reorganizing_in_one_file_is_not_a_false_cycle(client: TestClient, token_headers: dict[str, str]) -> None:
    """🔴 R1 ── A ── B、R2 是顶级。文件里 **A 写在前面**：A 挂到 B 下面，B 挪到 R2 下面。

    导完之后是 R2 ── B ── A，没有环。两种错法都会让这条红：
    - 预览按库里**现在**的树查环（那时 B 还是 A 的下级）→ 误报成环，这一行导不了
    - 提交按文件顺序写（先挂 A）→ 那一刻 B 还在 A 下面，service 的成环检查当场拒绝
    """
    make_dept(client, token_headers, 'R1')
    make_dept(client, token_headers, 'R2')
    make_dept(client, token_headers, 'MA', parent='R1')
    make_dept(client, token_headers, 'MB', parent='MA')

    raw = dept_rows(
        {'code': code('MA'), 'parent_code': code('MB')},
        {'code': code('MB'), 'parent_code': code('R2')},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw, 'update')
    assert (data['valid'], data['to_update']) == (2, 2), data['rows']

    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    assert body['code'] == 200, body
    assert body['data']['failed'] == [], body['data']
    tree = dept_tree(client, token_headers)
    assert tree[code('MB')]['parent_id'] == tree[code('R2')]['id']
    assert tree[code('MA')]['parent_id'] == tree[code('MB')]['id']


@pytest.mark.parametrize(
    ('mode', 'existing_action', 'new_action'),
    [('create', None, 'create'), ('upsert', 'update', 'create'), ('update', 'update', None)],
)
def test_modes_decide_what_an_existing_code_means(
    client: TestClient, token_headers: dict[str, str], mode: str, existing_action: str | None, new_action: str | None
) -> None:
    make_dept(client, token_headers, 'E')
    raw = dept_rows({'code': code('E'), 'name': f'导入{RUN}E2'}, {'code': code('N'), 'name': f'导入{RUN}N'})
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw, mode)
    existing, new = row_of(data, 'code', code('E')), row_of(data, 'code', code('N'))
    assert existing['action'] == existing_action, existing
    assert new['action'] == new_action, new
    # 不在这个模式里的那一行要说清楚为什么（「已存在」/「不存在」），不能只是静静地不导
    for row in (existing, new):
        if row['action'] is None:
            assert any(code('E') in msg or code('N') in msg for msg in row['errors']), row['errors']


def test_an_unsupported_mode_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    """用户导入只支持新增（同名就是另一个人，没有「更新」可言）"""
    body = preview(client, token_headers, USER_IMPORT, xlsx(['username', 'dept_code'], [['x', 'HQ']]), 'upsert')
    assert body['code'] != 200
    assert body['msg'] == t('error.import.mode_unsupported')


def test_blank_cells_keep_the_current_values_on_update(client: TestClient, token_headers: dict[str, str]) -> None:
    """只改一列的人不用把整行抄全 —— 而且文件里干脆没有的列（名称 / 上级）也一样是「不改」"""
    make_dept(client, token_headers, 'K', phone='13800138000', email='k@example.com', sort=3)
    raw = xlsx(['部门编码', '排序'], [[code('K'), 9]])
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw, 'update')
    assert data['valid'] == 1, data['rows']
    assert commit(client, token_headers, DEPT_IMPORT, data['import_token'])['data']['updated'] == [code('K')]

    node = dept_tree(client, token_headers)[code('K')]
    assert node['sort'] == 9  # 先断言「改了」：否则「什么都没写」也满足下面几条
    assert (node['name'], node['phone'], node['email'], node['parent_id']) == (
        f'导入{RUN}K',
        '13800138000',
        'k@example.com',
        None,
    )


def test_a_new_department_needs_a_name(client: TestClient, token_headers: dict[str, str]) -> None:
    """名称不是整列必填（更新时空白表示不改），新建时才必填 —— 要报在那一行上，而且说的是列名"""
    data = ok_preview(client, token_headers, DEPT_IMPORT, dept_rows({'code': code('NONAME')}))
    row = row_of(data, 'code', code('NONAME'))
    assert row['ok'] is False
    # 断言到「必填」这句话：只看列名的话，pydantic 自己那条「输入应为有效的字符串」也是列名打头
    assert row['errors'] == ['部门名称：必填项不能为空'], row['errors']


def test_child_of_an_invalid_new_parent_is_flagged_in_preview(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    """上级是同一份文件里新建的、而那一行本身有问题：下级在预览里就要标出来，别等提交时才冒出来"""
    raw = dept_rows(
        {'code': code('BADP'), 'name': f'导入{RUN}BADP', 'email': 'not-an-email'},
        {'code': code('KID'), 'name': 'KID', 'parent_code': code('BADP')},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['valid'] == 0
    kid = row_of(data, 'code', code('KID'))
    assert any(code('BADP') in msg for msg in kid['errors']), kid['errors']


def test_siblings_with_the_same_name_in_one_file(client: TestClient, token_headers: dict[str, str]) -> None:
    """两行都是新的、库里查同级重名都通过 —— 只有文件里的行之间比才查得出来"""
    make_dept(client, token_headers, 'SP')
    raw = dept_rows(
        {'code': code('S1'), 'name': 'dup', 'parent_code': code('SP')},
        {'code': code('S2'), 'name': 'dup', 'parent_code': code('SP')},
        {'code': code('S3'), 'name': 'dup', 'parent_code': code('S1')},  # 不同上级下同名是合法的
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert row_of(data, 'code', code('S1'))['ok'] is True
    assert row_of(data, 'code', code('S3'))['ok'] is True
    s2 = row_of(data, 'code', code('S2'))
    assert s2['ok'] is False and any(str(row_of(data, 'code', code('S1'))['row_no']) in m for m in s2['errors'])


def test_leader_is_a_username_and_must_be_an_active_account(client: TestClient, token_headers: dict[str, str]) -> None:
    make_user('lead')
    make_user('gone', status=0)
    raw = dept_rows(
        {'code': code('L1'), 'name': f'导入{RUN}L1', 'leader': username('lead')},
        {'code': code('L2'), 'name': f'导入{RUN}L2', 'leader': username('gone')},
        {'code': code('L3'), 'name': f'导入{RUN}L3', 'leader': username('nobody')},
        # 上级是文件里新建的 L1：service 那道负责人校验要留到提交才跑，预览时只有导入自己这一道
        {'code': code('L4'), 'name': f'导入{RUN}L4', 'parent_code': code('L1'), 'leader': username('gone')},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert row_of(data, 'code', code('L1'))['ok'] is True
    for key in ('L2', 'L3', 'L4'):
        assert row_of(data, 'code', code(key))['ok'] is False, key

    commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    assert dept_tree(client, token_headers)[code('L1')]['leader_name'] == '导入lead'


def test_a_disabled_current_leader_does_not_block_reimporting(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    """负责人离职（账号停用）之后，把导出的表原样导回来不能卡在「负责人已停用」上 ——
    service 也是「换了负责人才校验」，否则这个部门就再也改不了"""
    leader_id = make_user('left')
    make_dept(client, token_headers, 'LD', leader_id=str(leader_id))
    asyncio.run(_sql('UPDATE sys_user SET status = 0 WHERE id = :pk', {'pk': leader_id}))

    raw = dept_rows({'code': code('LD'), 'leader': username('left'), 'sort': 5})
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw, 'update')
    assert data['valid'] == 1, data['rows']
    # 对照：把他**新**设成另一个部门的负责人，照样要拦
    data = ok_preview(
        client, token_headers, DEPT_IMPORT, dept_rows({'code': code('LE'), 'name': 'x', 'leader': username('left')})
    )
    assert data['valid'] == 0


def test_status_accepts_the_word_in_either_language_or_the_value(
    client: TestClient, token_headers: dict[str, str]
) -> None:
    raw = dept_rows(
        {'code': code('ST1'), 'name': f'导入{RUN}ST1', 'status': '停用'},
        {'code': code('ST2'), 'name': f'导入{RUN}ST2', 'status': 'Disabled'},
        {'code': code('ST3'), 'name': f'导入{RUN}ST3', 'status': '0'},
        {'code': code('ST4'), 'name': f'导入{RUN}ST4', 'status': '半停用'},
        # 下拉里是「正常 / 停用」，手打的常常是「启用 / 禁用」—— 目测时就是这么填的，被拒了
        {'code': code('ST5'), 'name': f'导入{RUN}ST5', 'status': '禁用'},
        {'code': code('ST6'), 'name': f'导入{RUN}ST6', 'status': '启用'},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['valid'] == 5, data['rows']
    assert any('状态' in msg for msg in row_of(data, 'code', code('ST4'))['errors'])
    commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    tree = dept_tree(client, token_headers)
    assert [tree[code(key)]['status'] for key in ('ST1', 'ST2', 'ST3', 'ST5', 'ST6')] == [0, 0, 0, 0, 1]


def test_an_exported_file_imports_back_unchanged(client: TestClient, token_headers: dict[str, str]) -> None:
    """「导出的文件改完能原样导回来」：同一套列名，多出来的导出专有列（ID、创建时间……）只是被忽略"""
    make_dept(client, token_headers, 'RT')
    make_dept(client, token_headers, 'RTC', parent='RT', phone='13800138001')

    job = client.post(f'{DEPTS}/export', headers=token_headers, json={'filters': {'code': f'IMP{RUN}RT'}}).json()
    assert job['code'] == 200, job
    raw = client.get(f'/sys/exports/{job["data"]["id"]}/download', headers=token_headers).content

    data = ok_preview(client, token_headers, DEPT_IMPORT, raw, 'upsert')
    assert (data['total'], data['valid'], data['to_update']) == (2, 2, 2), data['rows']
    assert '部门 ID' in data['ignored_headers']
    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    assert sorted(body['data']['updated']) == [code('RT'), code('RTC')]
    tree = dept_tree(client, token_headers)
    assert tree[code('RTC')]['parent_id'] == tree[code('RT')]['id']
    assert tree[code('RTC')]['phone'] == '13800138001'


# ─────────────────────────────── 数据范围 ───────────────────────────────


@pytest.fixture
def scoped_importer(client: TestClient, token_headers: dict[str, str]) -> Iterator[dict[str, str]]:
    """一个有「部门导入」权限、但只看得见部门 IN 的人。返回他的请求头"""
    inside = make_dept(client, token_headers, 'IN')
    make_dept(client, token_headers, 'OUT')
    sid = snowflake.generate
    now = timezone.now()
    rule_id, scope_id, role_id = sid(), sid(), sid()
    user_id = make_user('scoped', dept_id=int(inside))
    rows = {
        DataRule.__table__: [
            {
                'id': rule_id,
                'name': f'IMP{RUN}-rule',
                # 按编码不按 ID：`id` 在 `DATA_PERMISSION_COLUMN_EXCLUDE` 里，按它写的规则会被整条跳过 ——
                # 角色开了范围过滤却没有一条生效的规则，就是 fail-closed：什么都看不见
                'model': 'Dept',
                'column': 'code',
                'operator': int(Op.AND),
                'expression': int(Expr.eq),
                'value': code('IN'),
                'deleted': 0,
                'deleted_time': None,
                'created_time': now,
                'updated_time': None,
            }
        ],
        DataScope.__table__: [
            {
                'id': scope_id,
                'name': f'IMP{RUN}-scope',
                'status': 1,
                'deleted': 0,
                'deleted_time': None,
                'created_time': now,
                'updated_time': None,
            }
        ],
        Role.__table__: [
            {
                'id': role_id,
                'code': f'IMP{RUN}ROLE',
                'name': f'IMP{RUN}-role',
                'status': 1,
                'is_filter_scopes': True,
                'remark': None,
                'deleted': 0,
                'deleted_time': None,
                'created_time': now,
                'updated_time': None,
            }
        ],
        data_scope_rule: [{'id': sid(), 'data_scope_id': scope_id, 'data_rule_id': rule_id}],
        role_data_scope: [{'id': sid(), 'role_id': role_id, 'data_scope_id': scope_id}],
        role_menu: [
            {'id': sid(), 'role_id': role_id, 'menu_id': menu} for menu in (SYS_DEPT_MENU_ID, IMPORT_SYS_DEPT_MENU_ID)
        ],
        user_role: [{'id': sid(), 'user_id': user_id, 'role_id': role_id}],
    }

    async def _write(*, remove: bool) -> None:
        engine = create_database_async_engine(get_database_url(unittest=True))
        try:
            async with engine.begin() as conn:
                for table, values in reversed(rows.items()) if remove else rows.items():
                    if remove:
                        await conn.execute(delete(table).where(table.c.id.in_([v['id'] for v in values])))
                    else:
                        await conn.execute(insert(table), values)
        finally:
            await engine.dispose()

    asyncio.run(_write(remove=False))
    try:
        res = client.post('/auth/login/swagger', params={'username': username('scoped'), 'password': PASSWORD})
        assert res.status_code == 200, res.text
        body = res.json()
        yield {'Authorization': f'{body["token_type"]} {body["access_token"]}'}
    finally:
        asyncio.run(_write(remove=True))


def test_an_importer_can_only_touch_departments_they_can_see(
    client: TestClient, token_headers: dict[str, str], scoped_importer: dict[str, str]
) -> None:
    """🔴 导入是批量的「新增 / 修改」，数据范围必须和表单同一个口径：看不见的上级不能挂，看不见的部门不能改。

    「先断言有」：挂在看得见的 IN 下面那一行必须能导 —— 否则一个什么都拒绝的实现也能让后两条绿
    """
    raw = dept_rows(
        {'code': code('NEWIN'), 'name': f'导入{RUN}NEWIN', 'parent_code': code('IN')},
        {'code': code('NEWOUT'), 'name': f'导入{RUN}NEWOUT', 'parent_code': code('OUT')},
        # 上级是同一份文件里新建的 NEWIN：这种行 service 的校验要留到提交才跑（库里还没有上级），
        # 预览时「看不看得见 OUT」只有导入自己那一道 —— 漏了它，预览会说「能导」
        {'code': code('OUT'), 'name': f'导入{RUN}被改名', 'parent_code': code('NEWIN')},
    )
    data = ok_preview(client, scoped_importer, DEPT_IMPORT, raw, 'upsert')
    assert row_of(data, 'code', code('NEWIN'))['ok'] is True, [
        (row['values']['code'], row['errors']) for row in data['rows']
    ]
    assert row_of(data, 'code', code('NEWOUT'))['ok'] is False
    assert row_of(data, 'code', code('OUT'))['ok'] is False

    body = commit(client, scoped_importer, DEPT_IMPORT, data['import_token'])
    assert body['data']['created'] == [code('NEWIN')], body
    assert dept_tree(client, token_headers)[code('OUT')]['name'] == f'导入{RUN}OUT'


# ═══════════════════════════════ 框架 ═══════════════════════════════


def test_every_import_is_translated_in_both_languages() -> None:
    """模板表头、说明页、取值文字按导入人的语言出：一边缺键，那个语言的人拿到的就是一串键名"""
    specs = import_specs()
    assert {'sys.user', 'sys.dept', 'dict.data'} <= set(specs)  # 先断言「有」
    for language in ('zh-CN', 'en-US'):
        with i18n.use_language(language):
            for spec in specs.values():
                keys = [spec.label, *spec.notes]
                for column in spec.columns:
                    keys += [column.label, *([column.note] if column.note else []), *(column.choices or {}).values()]
                missing = [key for key in keys if t(key) == key]
                assert not missing, f'{language} 缺 {spec.resource}: {missing}'


@pytest.mark.parametrize('language', ['zh-CN', 'en-US'])
@pytest.mark.parametrize(
    ('base', 'resource'), [(USER_IMPORT, 'sys.user'), (DEPT_IMPORT, 'sys.dept'), (DICT_IMPORT, 'dict.data')]
)
def test_every_template_parses_back_in_both_languages(
    client: TestClient, token_headers: dict[str, str], base: str, resource: str, language: str
) -> None:
    """下发的模板（任一语言）原样交回解析器，表头一个都不能落空 —— 模板和解析器分叉是无声的"""
    res = client.get(f'{base}/template', headers={**token_headers, 'Accept-Language': language})
    assert res.status_code == 200, res.text
    spec = import_specs()[resource]
    wb = load_workbook(io.BytesIO(res.content))
    header = next(wb.worksheets[0].iter_rows(max_row=1, values_only=True))
    with i18n.use_language(language):
        assert list(header) == [t(column.label) for column in spec.columns]

    wb.worksheets[0].append([f'v{index}' for index in range(len(spec.columns))])
    filled = io.BytesIO()
    wb.save(filled)
    parsed = parse_table(filled.getvalue(), _table_columns(spec))
    assert parsed.ignored_headers == []
    # 每一列都读到了自己那一格：表头认错列（两列互换）也是「全认得」，但值进了别的列
    assert parsed.rows[0].values == {column.key: f'v{index}' for index, column in enumerate(spec.columns)}


def test_headers_in_the_other_language_are_recognized(client: TestClient, token_headers: dict[str, str]) -> None:
    """英文界面导出的表交给中文界面的同事导入：表头按任一语言都认"""
    raw = xlsx(['Department code', 'Department name'], [[code('EN'), f'导入{RUN}EN']])
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['ignored_headers'] == []
    assert data['valid'] == 1, data['rows']


def test_a_token_works_only_for_its_own_resource_and_owner(
    client: TestClient, token_headers: dict[str, str], scoped_importer: dict[str, str]
) -> None:
    data = ok_preview(client, token_headers, DEPT_IMPORT, dept_rows({'code': code('TK'), 'name': f'导入{RUN}TK'}))
    token = data['import_token']
    # 拿去别的资源提交：拒绝，而且**不消耗**这个令牌（否则拼错接口的人连自己的导入都提交不了）
    assert commit(client, token_headers, DICT_IMPORT, token)['code'] != 200
    # 别人拿去提交：同样拒绝、同样不消耗
    assert commit(client, scoped_importer, DEPT_IMPORT, token)['code'] != 200
    body = commit(client, token_headers, DEPT_IMPORT, token)
    assert body['code'] == 200 and body['data']['created'] == [code('TK')], body
    # 用过一次就没了（GETDEL）：重放不会建第二遍
    assert commit(client, token_headers, DEPT_IMPORT, token)['code'] != 200


def test_excluding_every_row_is_rejected(client: TestClient, token_headers: dict[str, str]) -> None:
    data = ok_preview(client, token_headers, DEPT_IMPORT, dept_rows({'code': code('EX'), 'name': f'导入{RUN}EX'}))
    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'], exclude_rows=[2])
    assert body['code'] != 200
    assert code('EX') not in dept_tree(client, token_headers)


def test_the_error_report_has_the_original_cells_and_the_reasons(
    client: TestClient, token_headers: dict[str, str], scoped_importer: dict[str, str]
) -> None:
    """有问题的行原样写回一张表（用户改完能直接再导），只给发起人"""
    raw = dept_rows(
        {'code': code('GOOD'), 'name': f'导入{RUN}GOOD'},
        {'code': code('BAD'), 'name': f'导入{RUN}BAD', 'email': 'not-an-email', 'phone': '13800138009'},
    )
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['has_report'] is True
    path = f'/sys/imports/{data["import_token"]}/report'
    assert client.get(path, headers=scoped_importer).status_code == 404

    res = client.get(path, headers=token_headers)
    assert res.status_code == 200, res.text
    rows = list(load_workbook(io.BytesIO(res.content)).worksheets[0].iter_rows(values_only=True))
    assert list(rows[0]) == [*DEPT_HEADER, '原行号', '错误原因']
    assert len(rows) == 2  # 只有有问题的那一行
    bad = rows[1]
    assert (bad[0], bad[4], bad[5]) == (code('BAD'), '13800138009', 'not-an-email')  # 原样，包括错的那一格
    assert bad[8] == 3 and '邮箱' in bad[9]

    # 改好那一行，把报告原样交回去：多出来的两列只是被忽略
    wb = load_workbook(io.BytesIO(res.content))
    wb.worksheets[0]['F2'] = 'ok@example.com'
    fixed = io.BytesIO()
    wb.save(fixed)
    again = ok_preview(client, token_headers, DEPT_IMPORT, fixed.getvalue())
    assert again['valid'] == 1, again['rows']
    assert again['ignored_headers'] == ['原行号', '错误原因']


def test_cell_errors_read_like_sentences_not_templates(client: TestClient, token_headers: dict[str, str]) -> None:
    """🔴 邮箱格式错时 pydantic 给的是 `{reason}`，而译文模板要 `{error}` —— 原来预览里显示的是
    「邮箱：值错误，{error}」。测试全绿，是看截图才发现的：原来的断言只查了「邮箱」两个字在不在"""
    raw = dept_rows({'code': code('MSG'), 'name': f'导入{RUN}MSG', 'email': 'not-an-email', 'phone': '123'})
    errors = row_of(ok_preview(client, token_headers, DEPT_IMPORT, raw), 'code', code('MSG'))['errors']
    assert sorted(errors) == ['手机号：格式不对，填法见模板的「说明」页', '邮箱：不是有效的邮箱地址'], errors
    # 没填上的占位符（`{error}` 这种）一个都不能有；正则里的 `{9}` 也不该出现在给人看的话里
    assert not any(re.search(r'\{\w*\}', msg) for msg in errors), errors


def test_untranslatable_errors_fall_back_to_pydantic_not_to_the_template() -> None:
    """译文要的占位符这条报错没给（`value_error` 的模板要 `{error}`，这条只给了 `reason`）：
    用 pydantic 自己那句，不把「值错误，{error}」原样摊给用户。邮箱那一种单独给了人话，
    这条守的是剩下那些没单独处理的"""
    from backend.app.admin.service.import_service import _cell_msg

    error = {'type': 'value_error', 'msg': 'Value error, something odd', 'ctx': {'reason': 'odd'}}
    assert _cell_msg(error) == 'Value error, something odd'
    # 对照：占位符给全了就用译文（否则「永远回落英文」的实现也能让上一条绿）
    assert _cell_msg({'type': 'value_error', 'msg': 'x', 'ctx': {'error': ValueError('坏了')}}) == '值错误，坏了'


def test_the_error_report_is_in_the_language_of_the_preview(client: TestClient, token_headers: dict[str, str]) -> None:
    """报告在线程里渲染、而且可能换一个语言的会话来下载 —— 两种情况都要按**预览时**那个人的语言出。
    默认语言是中文，所以只有拿英文测才分得出「按了预览的语言」和「什么都没按、碰巧是中文」"""
    en = {**token_headers, 'Accept-Language': 'en-US'}
    raw = dept_rows({'code': code('ENR'), 'name': f'导入{RUN}ENR', 'email': 'not-an-email'})
    data = ok_preview(client, en, DEPT_IMPORT, raw)
    assert data['has_report'] is True
    assert data['rows'][0]['errors'][0].startswith('Email: '), data['rows'][0]['errors']

    res = client.get(f'/sys/imports/{data["import_token"]}/report', headers=token_headers)
    assert res.status_code == 200, res.text
    header = next(load_workbook(io.BytesIO(res.content)).worksheets[0].iter_rows(max_row=1, values_only=True))
    assert list(header[-2:]) == ['Original row', 'Errors']
    assert header[0] == 'Department code'


@pytest.mark.parametrize('atomic', [True, False])
def test_a_row_that_breaks_between_preview_and_commit(
    client: TestClient,
    token_headers: dict[str, str],
    atomic: bool,  # ruff:ignore[boolean-type-hint-positional-argument]
) -> None:
    """预览之后、提交之前，别人抢先建了同一个编码。提交时那一行要报失败（不是 500）；
    `atomic=True` 时**一行都不写**，否则能导的照常导"""
    raw = dept_rows({'code': code('AT1'), 'name': f'导入{RUN}AT1'}, {'code': code('AT2'), 'name': f'导入{RUN}AT2'})
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    assert data['valid'] == 2
    make_dept(client, token_headers, 'AT2')

    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'], atomic=atomic)
    assert body['code'] == 200, body
    assert [item['row_no'] for item in body['data']['failed']] == [3]
    assert body['data']['committed'] is not atomic
    assert (code('AT1') in dept_tree(client, token_headers)) is not atomic
    assert body['data']['has_report'] is True


@pytest.mark.parametrize('atomic', [True, False])
def test_a_write_that_fails_after_other_rows_were_written(
    client: TestClient,
    token_headers: dict[str, str],
    monkeypatch: pytest.MonkeyPatch,
    atomic: bool,  # ruff:ignore[boolean-type-hint-positional-argument]
) -> None:
    """校验都过了、写到第二行时库报错（并发、约束、连接抖动）。第一行**已经写进去了**：
    `atomic=True` 要把它一起撤掉，否则「全有或全无」只是对校验失败成立"""
    from sqlalchemy.exc import OperationalError

    from backend.app.admin.dataio import dept_service

    original = dept_service.create

    async def flaky_create(*, db: Any, obj: Any) -> None:
        if obj.code == code('W2'):
            raise OperationalError('INSERT', {}, Exception('simulated'))
        await original(db=db, obj=obj)

    monkeypatch.setattr(dept_service, 'create', flaky_create)
    raw = dept_rows({'code': code('W1'), 'name': f'导入{RUN}W1'}, {'code': code('W2'), 'name': f'导入{RUN}W2'})
    data = ok_preview(client, token_headers, DEPT_IMPORT, raw)
    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'], atomic=atomic)
    assert body['code'] == 200, body
    assert [item['row_no'] for item in body['data']['failed']] == [3]
    tree = dept_tree(client, token_headers)
    assert code('W2') not in tree
    assert (code('W1') in tree) is not atomic


def test_upsert_follows_the_database_at_commit_time(client: TestClient, token_headers: dict[str, str]) -> None:
    """「新增并更新」：预览时还不存在的编码，提交前被别人建了 —— 提交时要按**那一刻**的库判成更新，
    而不是拿着预览时的「新增」去撞「编码已存在」（提交前那一遍重新校验就是为这个）"""
    data = ok_preview(
        client, token_headers, DEPT_IMPORT, dept_rows({'code': code('UP'), 'name': f'导入{RUN}UP2'}), 'upsert'
    )
    assert data['to_create'] == 1
    make_dept(client, token_headers, 'UP')

    body = commit(client, token_headers, DEPT_IMPORT, data['import_token'])
    assert body['code'] == 200, body
    assert (body['data']['created'], body['data']['updated'], body['data']['failed']) == ([], [code('UP')], [])
    assert dept_tree(client, token_headers)[code('UP')]['name'] == f'导入{RUN}UP2'


# ═══════════════════════════════ 字典数据 ═══════════════════════════════

DICT_HEADER = ['字典类型编码', '标签', '键值', '排序', '状态', '颜色', '备注']


@pytest.fixture
def dict_type(client: TestClient, token_headers: dict[str, str]) -> str:
    """一个字典类型，下面已经有一项 `a = 1`"""
    type_code = f'imp{RUN}_type'
    res = client.post('/sys/dict-types', headers=token_headers, json={'name': f'导入{RUN}', 'code': type_code}).json()
    assert res['code'] == 200, res
    types = client.get('/sys/dict-types', headers=token_headers, params={'code': type_code}).json()['data']['items']
    body = {'type_id': types[0]['id'], 'label': 'a', 'value': '1', 'sort': 0, 'status': 1}
    res = client.post('/sys/dict-datas', headers=token_headers, json=body).json()
    assert res['code'] == 200, res
    return type_code


def dict_values(client: TestClient, headers: dict[str, str], type_code: str) -> dict[str, str]:
    """经**带缓存**的那个接口读（`get_by_type_code` 挂着 `@cached`）"""
    res = client.get(f'/sys/dict-datas/type-codes/{type_code}', headers=headers).json()
    assert res['code'] == 200, res
    return {item['label']: item['value'] for item in res['data']}


def test_dict_import_creates_and_updates_by_type_and_label(
    client: TestClient, token_headers: dict[str, str], dict_type: str
) -> None:
    assert dict_values(client, token_headers, dict_type) == {'a': '1'}  # 顺带把缓存焐热
    raw = xlsx(
        DICT_HEADER,
        [[dict_type, 'a', '2', None, None, None, None], [dict_type, 'b', '3'], ['no_such_type', 'c', '4']],
    )
    data = ok_preview(client, token_headers, DICT_IMPORT, raw, 'upsert')
    assert (data['to_update'], data['to_create'], data['valid']) == (1, 1, 2), data['rows']
    assert any('no_such_type' in msg for msg in row_of(data, 'label', 'c')['errors'])

    body = commit(client, token_headers, DICT_IMPORT, data['import_token'])
    assert body['code'] == 200, body
    # 🔴 读的是带缓存的接口：提交之后没清缓存的话，这里还是导入之前的 `{'a': '1'}`
    assert dict_values(client, token_headers, dict_type) == {'a': '2', 'b': '3'}


def test_dict_value_is_required_only_when_creating(
    client: TestClient, token_headers: dict[str, str], dict_type: str
) -> None:
    raw = xlsx(DICT_HEADER, [[dict_type, 'a', None, 7], [dict_type, 'new']])
    data = ok_preview(client, token_headers, DICT_IMPORT, raw, 'upsert')
    assert row_of(data, 'label', 'a')['ok'] is True  # 更新：空白 = 不改
    assert row_of(data, 'label', 'new')['errors'] == ['键值：必填项不能为空']
    commit(client, token_headers, DICT_IMPORT, data['import_token'])
    assert dict_values(client, token_headers, dict_type) == {'a': '1'}
