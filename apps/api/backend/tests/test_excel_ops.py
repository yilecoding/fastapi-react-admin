"""表格解析层（`utils/excel_ops.py`）的测试。

这份文件里几乎每一条都对应一种**「解析成功、有数据、就是错的」**的情况 ——
纯函数、不连库，跑起来是毫秒级的，所以没有理由不覆盖全。

⚠️ 断言写的是**行为**不是实现：比如「公式格不会变成 None」而不是
「`data_only` 传了 False」。换解析库时这批测试应该还能用。
"""

import asyncio
import datetime as dt
import io
import re
import zipfile

from decimal import Decimal

import pytest

from fastapi import UploadFile
from openpyxl import Workbook

from backend.common.exception import errors
from backend.core.conf import settings
from backend.utils.excel_ops import (
    TableColumn,
    _CellError,
    _normalize_value,
    build_template,
    parse_table,
    read_table_upload,
)

COLUMNS = [
    TableColumn(key='username', required=True),
    TableColumn(key='phone'),
    TableColumn(key='dept_code', required=True, aliases=('dept',)),
]


def make_xlsx(rows: list[list], *, sheets: dict[str, list[list]] | None = None, active: int = 0) -> bytes:
    """造一份 xlsx。`rows` 是第一个工作表的内容（含表头）"""
    wb = Workbook()
    ws = wb.active
    ws.title = 'data'
    for row in rows:
        ws.append(row)
    for title, content in (sheets or {}).items():
        extra = wb.create_sheet(title)
        for row in content:
            extra.append(row)
    wb.active = active
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def values_of(raw: bytes, columns=COLUMNS, **kwargs) -> list[dict]:
    return [row.values for row in parse_table(raw, columns, **kwargs).rows]


# --------------------------------------------------------------------------- 选表


def test_reads_the_first_sheet_not_the_one_the_user_left_selected():
    """🔴 `wb.active` 是「保存时停在哪个页签」，不是第一个表。

    实测：第一个表叫 data、第二个叫「临时草稿」，保存时停在第二个页签，
    `wb.active.title` 就是「临时草稿」—— 解析成功、有数据、读的是草稿。
    这条是本模块最需要守住的一条，因为它没有任何报错。
    """
    raw = make_xlsx(
        [['username', 'dept_code'], ['real', 'DEPT_0001']],
        sheets={'临时草稿': [['username', 'dept_code'], ['GARBAGE', 'DEPT_9999']]},
        active=1,
    )
    assert values_of(raw)[0]['username'] == 'real'


def test_sheet_can_be_picked_by_name():
    raw = make_xlsx(
        [['username', 'dept_code'], ['first', 'DEPT_0001']],
        sheets={'别的': [['username', 'dept_code'], ['second', 'DEPT_0002']]},
    )
    assert values_of(raw, sheet_name='别的')[0]['username'] == 'second'


def test_unknown_sheet_name_is_a_clear_error():
    raw = make_xlsx([['username', 'dept_code'], ['a', 'DEPT_0001']])
    with pytest.raises(errors.RequestError, match='工作表'):
        parse_table(raw, COLUMNS, sheet_name='不存在')


# --------------------------------------------------------------------------- 表头


def test_headers_are_matched_by_text_not_by_position():
    """用户会插列、删列、调顺序，这三件事都不会让按位置解析的代码报错"""
    raw = make_xlsx([
        ['dept_code', '备注', 'username', 'phone'],
        ['DEPT_0001', '随手写的', 'zhangsan', '13800138000'],
    ])
    assert values_of(raw)[0] == {'username': 'zhangsan', 'phone': '13800138000', 'dept_code': 'DEPT_0001'}


def test_missing_required_header_rejects_the_whole_file():
    raw = make_xlsx([['username', 'phone'], ['zhangsan', '13800138000']])
    with pytest.raises(errors.RequestError, match='dept_code'):
        parse_table(raw, COLUMNS)


def test_duplicated_header_rejects_the_whole_file():
    """同一列出现两次时后一列会静静盖掉前一列 —— 宁可整份拒掉"""
    raw = make_xlsx([['username', 'dept_code', 'username'], ['a', 'DEPT_0001', 'b']])
    with pytest.raises(errors.RequestError, match='username'):
        parse_table(raw, COLUMNS)


def test_unrecognized_header_is_reported_not_silently_dropped():
    """`phone` 拼成 `phone_number`：必填列不受影响，但那一列的数据全丢了。

    非必填列的拼写错误是**纯静默**的 —— 没有它整份文件照样解析成功。
    所以必须把「没认出来的列」回报给调用方去提示用户。
    """
    raw = make_xlsx([
        ['username', 'dept_code', 'phone_number'],
        ['zhangsan', 'DEPT_0001', '13800138000'],
    ])
    result = parse_table(raw, COLUMNS)
    assert result.ignored_headers == ['phone_number']
    assert result.rows[0].values['phone'] is None


def test_header_alias_and_case_are_tolerated():
    raw = make_xlsx([['  USERNAME ', 'Dept'], ['zhangsan', 'DEPT_0001']])
    assert values_of(raw)[0]['dept_code'] == 'DEPT_0001'


# --------------------------------------------------------------------------- 单元格归一


def test_phone_typed_as_a_number_comes_back_as_plain_digits():
    """用户直接打手机号时 Excel 存的是**数字**不是文本，读回来必须还是那串数字。

    ⚠️ 这条走的是 int 分支 —— 实测 openpyxl 会把整数值的 float 收成 int，
    所以经过它的路径**测不到** `.0` 尾巴那个分支。那个分支由下面的
    `test_normalize_value_truth_table` 直接单测。
    """
    raw = make_xlsx([['username', 'phone', 'dept_code'], ['zhangsan', 13800138000, 'DEPT_0001']])
    assert values_of(raw)[0]['phone'] == '13800138000'


def test_number_beyond_the_safe_range_is_rejected_not_silently_rounded():
    """🔴 19 位数字在 Excel 里已经塌成 float 了（Excel 用 IEEE double 存数值）。

    `int(1.234567890123457e+18)` 会给出 `1234567890123456768` —— 一个
    **看起来完全正常、但和用户填的不是同一个数**的整数。同硬纪律 6。
    """
    raw = make_xlsx([['username', 'phone', 'dept_code'], ['zhangsan', 1234567890123456789, 'DEPT_0001']])
    result = parse_table(raw, COLUMNS)
    assert result.rows[0].values['phone'] is None
    assert [(i.row_no, i.column) for i in result.issues] == [(2, 'phone')]
    assert '文本' in result.issues[0].msg


def test_formula_becomes_an_issue_instead_of_an_empty_value():
    """🔴 按缓存值读公式，在程序生成的表格上会得到 `None` —— 和空格无法区分。

    用户拿公式拼用户名时，那一行会静静变成「用户名没填」。
    """
    raw = make_xlsx([['username', 'dept_code'], ['=CONCAT("a","b")', 'DEPT_0001']])
    result = parse_table(raw, COLUMNS)
    assert [(i.row_no, i.column) for i in result.issues] == [(2, 'username')]
    assert '公式' in result.issues[0].msg


def test_leading_and_trailing_whitespace_is_stripped():
    """粘贴过来的值经常带前后空格。不 strip 的话「张三 」和「张三」是两个不同的
    用户名，而肉眼一模一样。

    首尾的**不换行空格**（U+00A0）不用额外处理 —— 实测 `'\\xa0'.isspace()` 在
    Python 里是 `True`，`str.strip()` 自己就剥掉了。真正需要处理的是夹在中间的
    那种，见下一条。
    """
    raw = make_xlsx([['username', 'dept_code'], ['\u00a0 zhangsan  ', ' DEPT_0001 ']])
    assert values_of(raw)[0] == {'username': 'zhangsan', 'phone': None, 'dept_code': 'DEPT_0001'}


def test_nbsp_in_the_middle_of_a_value_is_normalized_to_a_plain_space():
    """🔴 夹在中间的 NBSP 是纯静默的：`'zhang\\xa0san'` 和 `'zhang san'` 肉眼
    完全一样，而 `==` 为 False。

    它躲得过 `strip()`（只管首尾），也躲得过肉眼复核 —— 用户比对 Excel 和
    列表页会觉得两边一模一样，然后困惑于「为什么说这个用户名不存在」。

    ⚠️ NBSP 一律写 `\\u00a0` 转义，不要直接敲一个进去：源码里的 NBSP 是不可见的，
    编辑器的「清理空白」或一次复制粘贴就能把它换成普通空格，而那之后这条测试
    照样是绿的 —— 它只是不再测 NBSP 了。
    """
    raw = make_xlsx([['username', 'dept_code'], ['zhang\u00a0san', 'DEPT_0001']])
    assert values_of(raw)[0]['username'] == 'zhang san'


def test_blank_looking_cell_is_none_not_empty_string():
    """只有空格的格子等同于没填 —— 否则必填校验会被一个空格骗过去"""
    raw = make_xlsx([['username', 'phone', 'dept_code'], ['zhangsan', '   ', 'DEPT_0001']])
    assert values_of(raw)[0]['phone'] is None


def test_bool_is_not_stringified_as_one_or_zero():
    """`isinstance(True, int)` 在 Python 里是 True，判断顺序反了 True 会变成 '1'"""
    raw = make_xlsx([['username', 'phone', 'dept_code'], ['zhangsan', True, 'DEPT_0001']])
    assert values_of(raw)[0]['phone'] == 'true'


def test_integral_float_keeps_no_decimal_tail():
    raw = make_xlsx([['username', 'phone', 'dept_code'], ['zhangsan', 1.0, 'DEPT_0001']])
    assert values_of(raw)[0]['phone'] == '1'


# --------------------------------------------------------------------------- 真值表（不经过 openpyxl）


@pytest.mark.parametrize(
    ('value', 'expected'),
    [
        (None, None),
        ('', None),
        ('   ', None),
        ('  admin  ', 'admin'),
        (True, 'true'),
        (False, 'false'),
        (0, '0'),
        (13800138000, '13800138000'),
        # 🔴 走这里的是「换了解析库」的情况：python-calamine 对同一个格子给的是
        # float（实测），裸 str() 会得到 '13800138000.0' 并一路活进数据库
        (13800138000.0, '13800138000'),
        (1.0, '1'),
        (1.5, '1.5'),
        (Decimal('1.50'), '1.50'),
        # ⚠️ 这里的 naive datetime 是**刻意**的，不要按 lint 的建议补 tzinfo：
        # Excel 的日期单元格本来就不带时区（它只存一个从 1900 起的序列号），
        # 补一个时区进去就不再是这一层真实会收到的输入了。要不要按某个时区
        # 解释它是**业务层**的决定，解析层原样交出去
        (dt.datetime(2026, 9, 7, 10, 30), '2026-09-07T10:30:00'),  # ruff: ignore[call-datetime-without-tzinfo]
        (dt.date(2026, 9, 7), '2026-09-07'),
        (dt.time(10, 30), '10:30:00'),
    ],
)
def test_normalize_value_truth_table(value, expected):
    """直接单测归一函数。

    `_normalize_value` 从 `_normalize` 里拆出来就是为了这个：有几个分支
    **经过 openpyxl 是走不到的**（它会先替我们把 float 收成 int），
    只能在这里钉住。走不到不等于不需要 —— 它们防的是换库。
    """
    assert _normalize_value(value) == expected


@pytest.mark.parametrize('value', [2**53, -(2**53), 1234567890123456789, 1.2345678901234568e18])
def test_normalize_value_rejects_numbers_beyond_the_safe_range(value):
    with pytest.raises(_CellError):
        _normalize_value(value)


def test_normalize_value_refuses_to_guess_at_unknown_types():
    """不要 `str(value)` 兜底 —— 那正是「解析成功但值是错的」的来源"""
    with pytest.raises(_CellError):
        _normalize_value(object())


# --------------------------------------------------------------------------- 行


def test_fully_empty_row_is_skipped_but_partially_empty_row_is_reported():
    """尾部空行极常见，跳过；但半空的行必须报出来。

    两者混为一谈的后果是「用户名忘填了」被当成「这行不存在」而消失 ——
    用户数了数导入结果，少了一个人，而界面上没有任何错误。
    """
    raw = make_xlsx([
        ['username', 'phone', 'dept_code'],
        ['zhangsan', '13800138000', 'DEPT_0001'],
        [None, None, None],
        [None, '13900139000', 'DEPT_0001'],
    ])
    result = parse_table(raw, COLUMNS)
    assert result.empty_rows == 1
    assert [row.row_no for row in result.rows] == [2, 4]
    assert [(i.row_no, i.column) for i in result.issues] == [(4, 'username')]


def test_row_numbers_are_the_real_excel_row_numbers():
    """报给用户的行号必须能在 Excel 里直接定位，所以从 2 起（1 是表头）"""
    raw = make_xlsx([['username', 'dept_code'], ['a', 'DEPT_0001'], ['b', 'DEPT_0002']])
    assert [row.row_no for row in parse_table(raw, COLUMNS).rows] == [2, 3]


def test_row_cap_rejects_a_file_that_really_is_too_long():
    rows = [['username', 'dept_code']] + [[f'u{i}', 'DEPT_0001'] for i in range(5)]
    with pytest.raises(errors.RequestError, match='3'):
        parse_table(make_xlsx(rows), COLUMNS, max_rows=3)


def test_row_cap_does_not_trust_the_dimension_declared_in_the_file():
    """🔴 `ws.max_row` 是文件**自己声明**的，不是真实行数。

    实测：把 sheet XML 里的 `<dimension ref="A1:A3" />` 改成 `A1:A1048576`，
    `ws.max_row` 就返回 1048576，而真实数据只有 2 行。拿它做行数上限判断 =
    一个改一行 XML 就能触发的拒绝服务。
    """
    raw = make_xlsx([['username', 'dept_code'], ['a', 'DEPT_0001'], ['b', 'DEPT_0002']])
    spoofed = _spoof_dimension(raw)

    # 先确认伪造确实生效了 —— 否则这条测试什么都没测
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(spoofed), read_only=True)
    assert wb.worksheets[0].max_row == 1048576
    wb.close()

    result = parse_table(spoofed, COLUMNS, max_rows=10)
    assert [row.values['username'] for row in result.rows] == ['a', 'b']


def _spoof_dimension(raw: bytes) -> bytes:
    """把工作表声明的行数改成一百万行，内容不动"""
    zin = zipfile.ZipFile(io.BytesIO(raw))
    target = next(n for n in zin.namelist() if n.endswith('sheet1.xml'))
    xml = zin.read(target)
    # 只改行号，**保留原来的列范围** —— 连列一起改掉的话表头会只剩 A 列，
    # 测试会因为「缺少必需的列」而失败，看起来像是防御生效了，其实什么都没测到
    patched, n = re.subn(rb'(<dimension ref="[A-Z]+1:[A-Z]+)\d+(" ?/>)', rb'\g<1>1048576\g<2>', xml)
    assert n == 1, '没找到 dimension 标签，这条测试的前提不成立'

    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w') as zout:
        for name in zin.namelist():
            zout.writestr(name, patched if name == target else zin.read(name))
    return out.getvalue()


def test_a_file_with_only_a_header_row_is_an_error_not_an_empty_success():
    """只有表头 = 用户传错了文件。静静返回 0 行会让他以为「导入成功，0 个人」"""
    with pytest.raises(errors.RequestError, match='没有数据'):
        parse_table(make_xlsx([['username', 'dept_code']]), COLUMNS)


# --------------------------------------------------------------------------- 上传前的闸门


def upload(raw: bytes, filename: str = 'users.xlsx') -> bytes:
    return asyncio.run(read_table_upload(UploadFile(file=io.BytesIO(raw), filename=filename, size=len(raw))))


def test_xls_is_rejected_with_a_readable_message_not_a_zip_error():
    """openpyxl 读不了老的 .xls，不拦的话用户拿到的是「文件已损坏」"""
    with pytest.raises(errors.RequestError, match='xlsx'):
        upload(b'anything', filename='users.xls')


def test_csv_renamed_to_xlsx_is_rejected():
    with pytest.raises(errors.RequestError, match='损坏'):
        upload(b'username,dept_code\na,DEPT_0001\n')


def test_empty_file_is_rejected():
    with pytest.raises(errors.RequestError, match='损坏'):
        upload(b'')


def test_zip_bomb_is_rejected_before_openpyxl_ever_opens_it():
    """🔴 xlsx 就是 zip，所以它天然是解压炸弹载体。

    实测：199 KB 的上传能声明出 200 MB 的解压体积 —— 上传大小那道闸门
    对它完全无感。好在 zip 的中央目录里就写着解压后大小，不用真解压就能拦。
    """
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('xl/worksheets/sheet1.xml', b'0' * (settings.TABLE_IMPORT_UNCOMPRESSED_MAX + 1))
    raw = bomb.getvalue()

    assert len(raw) < settings.TABLE_IMPORT_SIZE_MAX, '前提：它压缩后小得能过大小闸门'
    with pytest.raises(errors.RequestError, match='解压'):
        upload(raw)


def test_oversized_upload_is_rejected():
    with pytest.raises(errors.RequestError, match='MB'):
        upload(b'x' * (settings.TABLE_IMPORT_SIZE_MAX + 1))


def test_a_good_file_passes_every_gate():
    raw = make_xlsx([['username', 'dept_code'], ['zhangsan', 'DEPT_0001']])
    assert upload(raw) == raw


# --------------------------------------------------------------------------- 模板


def test_template_round_trips_through_the_parser():
    """生成的模板必须能被自己的解析器读回来 —— 两边分叉是无声的"""
    template = build_template(COLUMNS, notes={'username': '登录名，必填'})
    wb_bytes = _fill(template, [['zhangsan', '13800138000', 'DEPT_0001']])
    assert values_of(wb_bytes)[0]['username'] == 'zhangsan'


def test_template_leaves_the_data_sheet_active():
    """说明页在第二个表。如果模板保存时停在说明页，用户随手另存就会踩上面那条坑"""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(build_template(COLUMNS)))
    assert wb.active.title == wb.worksheets[0].title == 'data'


def test_template_carries_the_notes_on_a_second_sheet():
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(build_template(COLUMNS, notes={'username': '登录名，必填'})))
    assert wb.sheetnames == ['data', 'README']
    assert '登录名，必填' in [row[2] for row in wb['README'].iter_rows(min_row=2, values_only=True)]


def _fill(template: bytes, rows: list[list]) -> bytes:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(template))
    for row in rows:
        wb.worksheets[0].append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
