"""
表格解析 —— 把上传的 xlsx 变成规整的行字典。

这一层**只管「读得对不对」，不管「业务上合不合法」**：它保证吐出来的每个值
要么是 `None`、要么是一个已经归一过的字符串，且任何「读不出来」的情况都会
变成一条**显式的** issue，而不是一个看起来很正常的空值。

🔴 **写这一层的全部理由就是最后半句。** Excel 有一大票「解析成功、有数据、
就是错的」的坑，实测确认过的在下面每个函数的注释里，一条对应一段防御。
业务层（比如批量导入用户）只声明列，不重复这些判断。
"""

import dataclasses
import datetime as dt
import io
import zipfile

from decimal import Decimal
from typing import Any

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.cell.read_only import EmptyCell

from backend.common.exception import errors
from backend.common.i18n import t
from backend.core.conf import settings

#: 只认这一种扩展名。openpyxl **读不了 `.xls`**（老 BIFF 格式，要另装 xlrd），
#: 硬拦在扩展名这一步是为了给一句人话，否则它会掉进 `zipfile.BadZipFile`
TABLE_EXT = 'xlsx'

#: 超过这个绝对值的整数，Excel 已经存不住了 —— 它用 IEEE double 存所有数值。
#: 同硬纪律 6 的 `Number.MAX_SAFE_INTEGER`，两边是同一个物理限制
_NUMBER_SAFE_MAX = 2**53 - 1


@dataclasses.dataclass(frozen=True)
class TableColumn:
    """一列的声明。业务层只需要给出这个，其余归一/校验由本模块负责"""

    #: 表头里的英文键，也是 `TableRow.values` 的键
    key: str
    #: 缺这一列时整份文件直接拒（区别于「某一行这一格为空」）
    required: bool = False
    #: 兼容的旧表头写法。比对前会统一 strip + 转小写，所以不用登记大小写变体
    aliases: tuple[str, ...] = ()


@dataclasses.dataclass
class RowIssue:
    """一处读不出来的地方。`column` 为 None 表示问题在整行/整份文件上"""

    #: Excel 里的**真实行号**（1 起，表头是 1）—— 报给用户的必须是这个，
    #: 不是「第几条数据」。用户要回 Excel 里定位，他看到的是行号
    row_no: int
    column: str | None
    msg: str


@dataclasses.dataclass
class TableRow:
    """一行数据。值要么是 `None`，要么是非空字符串 —— 不会有 `''`"""

    row_no: int
    values: dict[str, str | None]


@dataclasses.dataclass
class TableData:
    """一次解析的完整结果"""

    rows: list[TableRow]
    #: 单元格级问题。**有 issue 不等于 rows 为空** —— 好行照样返回，
    #: 由业务层决定是「整批拒绝」还是「跳过坏行」
    issues: list[RowIssue]
    #: 表头里出现、但没有被任何 `TableColumn` 认领的列。
    #: 必须回报给用户：`user_name` 拼错成这样时，非必填列会**静默丢数据**
    ignored_headers: list[str]
    #: 整行皆空而被跳过的行数。Excel 尾部空行极常见，不算错，但要能对账
    empty_rows: int


class _CellError(Exception):
    """单元格读不出来。只在本模块内部流转，出口一律转成 `RowIssue`"""

    def __init__(self, msg: str) -> None:
        self.msg = msg
        super().__init__(msg)


async def read_table_upload(file: UploadFile) -> bytes:
    """
    收下上传的表格文件并做进解析器**之前**的三道校验。

    ⚠️ 这条路**不走** `utils/file_ops.py` 那套上传管线：那套是为长期存档设计的
    （去重、按日期落盘、写 `sys_file`、下载鉴权），而导入用的表格是用完即丢的
    临时输入。硬套上去只会在文件管理页里堆一堆没人会点开的「导入模板.xlsx」。
    代价是**这里的三道校验得自己写**，`_upload_rules()` 那张白名单管不到我们。

    :param file: FastAPI 上传文件对象
    :return: 文件原始字节，供 `parse_table` 在内存里解析（全程不落盘）
    """
    filename = file.filename or ''
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext != TABLE_EXT:
        raise errors.RequestError(msg=t('error.table.not_xlsx', file_ext=ext or '-'))

    raw = await file.read()
    if len(raw) > settings.TABLE_IMPORT_SIZE_MAX:
        raise errors.RequestError(
            msg=t('error.table.size_exceeded', size_mb=settings.TABLE_IMPORT_SIZE_MAX // 1024 // 1024)
        )

    _verify_archive(raw)
    return raw


def _verify_archive(raw: bytes) -> None:
    """
    在把字节交给 openpyxl **之前**，先看一眼这个 zip 解压出来有多大。

    🔴 xlsx 就是一个 zip，所以它天然是个解压炸弹载体。实测：一个 **199 KB** 的
    上传能声明出 **200 MB** 的解压体积 —— 而 `TABLE_IMPORT_SIZE_MAX` 那道 5 MB
    的闸门对它完全无感（它压缩后确实只有 199 KB）。openpyxl 一旦开始读就是
    照着解压后的体积吃内存。

    好在**不用真解压就能拦**：zip 的中央目录里每一条都写着自己的
    `file_size`（解压后大小），`zipfile.infolist()` 打开时就读到了。
    """
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            uncompressed = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        # 走到这里通常是「后缀改成了 .xlsx 的别的东西」（csv / 空文件都会命中）。
        # 扩展名那道闸门只看名字，拦不住这个
        raise errors.RequestError(msg=t('error.table.corrupt')) from None

    if uncompressed > settings.TABLE_IMPORT_UNCOMPRESSED_MAX:
        raise errors.RequestError(
            msg=t('error.table.bomb', size_mb=settings.TABLE_IMPORT_UNCOMPRESSED_MAX // 1024 // 1024)
        )


def parse_table(
    raw: bytes,
    columns: list[TableColumn],
    *,
    sheet_name: str | None = None,
    max_rows: int | None = None,
) -> TableData:
    """
    解析一份 xlsx，返回规整的行。

    :param raw: 文件原始字节（走 `read_table_upload` 拿到）
    :param columns: 列声明
    :param sheet_name: 指定按名字取工作表；不传则取**第一个**（见下面的注释）
    :param max_rows: 数据行数上限，不传用 `TABLE_IMPORT_MAX_ROWS`
    :return:
    """
    limit = settings.TABLE_IMPORT_MAX_ROWS if max_rows is None else max_rows

    try:
        # read_only=True：实测 5000 行的峰值内存 0.9 MB，普通模式是 14 MB。
        #
        # data_only=**False** 是刻意的，反直觉但必须：data_only=True 时
        # openpyxl 把公式换成**缓存值**，而由程序生成的表格根本没有缓存值 ——
        # 公式格于是读出 `None`，和「这一格是空的」**完全无法区分**。
        # 用 False 读到的是公式原文，`data_type == 'f'` 能一眼认出来，
        # 于是「用户拿公式拼了用户名」这件事会变成一条显式 issue 而不是一个空值。
        # 代价只有 7%（实测 5000 行：values_only 186 ms → 取 cell 对象 199 ms）
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except Exception:
        # openpyxl 对坏文件抛的异常种类很杂（BadZipFile / KeyError / ValueError…），
        # 一律收敛成一句人话；真实原因对用户没有意义
        raise errors.RequestError(msg=t('error.table.corrupt')) from None

    try:
        ws = _pick_sheet(wb, sheet_name)
        rows_iter = ws.iter_rows()

        header_cells = next(rows_iter, None)
        if header_cells is None:
            raise errors.RequestError(msg=t('error.table.empty'))

        index_of, ignored_headers = _map_headers(header_cells, columns)

        rows: list[TableRow] = []
        issues: list[RowIssue] = []
        empty_rows = 0

        # 🔴 **不要用 `ws.max_row` 来判行数上限，它是文件自己声明的。**
        # 实测：把 sheet XML 里的 `<dimension ref="A1:A4"/>` 改成 `A1:A1048576`
        # 之后 `ws.max_row` 就返回 1048576，而真实数据只有 4 行。
        # 拿它做「太多行了，拒绝」的判断 = 一个改一行 XML 就能触发的拒绝服务；
        # 拿它预分配更糟。所以这里边迭代边数，超了立刻停 —— 也顺便意味着
        # 一份 100 万行的表**不会**被读完才报错
        for row_no, cells in enumerate(rows_iter, start=2):
            values: dict[str, str | None] = {}
            row_issues: list[RowIssue] = []

            for column in columns:
                idx = index_of.get(column.key)
                cell = cells[idx] if idx is not None and idx < len(cells) else None
                try:
                    values[column.key] = _normalize(cell)
                except _CellError as exc:
                    values[column.key] = None
                    row_issues.append(RowIssue(row_no=row_no, column=column.key, msg=exc.msg))

            # 整行皆空：Excel 尾部空行极常见（用户删过内容、或者选中过一片区域），
            # 静静跳过。**只有整行皆空才跳** —— 半空的行要走必填校验报出来，
            # 否则「用户名忘填了」会被当成「这行不存在」而消失
            if not row_issues and all(v is None for v in values.values()):
                empty_rows += 1
                continue

            # 已经报过问题的格子不再补一条「必填项为空」。
            #
            # 实测（就在这个模块自己的测试里发现的）：一个公式格会先被判成
            # 「不支持公式」、值回落成 None，然后又撞上必填校验 —— 用户看到
            # 同一个格子两条错，而第二条是**误导**的：他明明填了东西。
            # 一个格子只说一件事，说最靠近根因的那件
            rejected = {issue.column for issue in row_issues}
            row_issues.extend(
                RowIssue(row_no=row_no, column=column.key, msg=t('error.table.cell_required'))
                for column in columns
                if column.required and values[column.key] is None and column.key not in rejected
            )

            issues.extend(row_issues)
            rows.append(TableRow(row_no=row_no, values=values))

            if len(rows) > limit:
                raise errors.RequestError(msg=t('error.table.too_many_rows', max_rows=limit))

        if not rows:
            raise errors.RequestError(msg=t('error.table.empty'))

        return TableData(rows=rows, issues=issues, ignored_headers=ignored_headers, empty_rows=empty_rows)
    finally:
        wb.close()


def _pick_sheet(wb: Workbook, sheet_name: str | None) -> Any:
    """
    选工作表。

    🔴 **不要用 `wb.active`。** 它返回的是「文件保存时停留在哪个页签」，
    不是第一个工作表 —— 那个值是 workbook.xml 里的 `activeTab`，由用户
    保存前最后点了哪个页签决定。

    实测：一份第一个表叫「数据」、第二个表叫「临时草稿」的文件，保存时停在
    第二个页签，`wb.active.title` 就是「临时草稿」，读出来的是草稿里的内容。
    **解析成功、有数据、就是错的表**，而且没有任何一处会报错。
    """
    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            raise errors.RequestError(msg=t('error.table.sheet_not_found', sheet=sheet_name))
        return wb[sheet_name]

    if not wb.worksheets:
        raise errors.RequestError(msg=t('error.table.empty'))
    return wb.worksheets[0]


def _map_headers(header_cells: tuple[Any, ...], columns: list[TableColumn]) -> tuple[dict[str, int], list[str]]:
    """
    把表头文本映射成列下标。

    **按表头文本认列，不按位置认列** —— 用户会插列、删列、调顺序，而这三件事
    里没有一件会让按位置解析的代码报错，它只会开始读错列。

    :return: (列 key → 下标, 未被认领的表头文本)
    """
    lookup: dict[str, str] = {}
    for column in columns:
        for name in (column.key, *column.aliases):
            lookup[name.strip().lower()] = column.key

    index_of: dict[str, int] = {}
    ignored: list[str] = []

    for idx, cell in enumerate(header_cells):
        raw = cell.value if cell is not None else None
        text = str(raw).strip() if raw is not None else ''
        if not text:
            continue
        key = lookup.get(text.lower())
        if key is None:
            ignored.append(text)
            continue
        # 同一列出现两次：后一列会静静盖掉前一列。宁可整份拒掉，
        # 也不要让用户以为自己填的是生效的那一列
        if key in index_of:
            raise errors.RequestError(msg=t('error.table.header_duplicated', column=key))
        index_of[key] = idx

    missing = [c.key for c in columns if c.required and c.key not in index_of]
    if missing:
        raise errors.RequestError(msg=t('error.table.header_missing', columns=', '.join(missing)))

    return index_of, ignored


def _normalize(cell: Any) -> str | None:
    """
    把一个单元格归一成 `str | None`。空串一律收敛成 `None`。

    这个函数是整个模块的核心 —— 下面每一个分支都对应一种「不处理就会静默出错」
    的实测情况。
    """
    if cell is None or isinstance(cell, EmptyCell):
        return None

    # 公式：`data_type == 'f'` 时 `value` 是公式原文（因为上面用了 data_only=False）。
    # 显式拒绝而不是求值 —— 我们没有 Excel 的计算引擎，而「按缓存值读」在
    # 程序生成的表格上会得到 None，那正是要避免的那种静默
    if getattr(cell, 'data_type', None) == 'f':
        raise _CellError(t('error.table.cell_formula'))

    return _normalize_value(cell.value)


def _normalize_value(value: Any) -> str | None:
    """值层面的归一。拆出来是为了能脱离 openpyxl 单测这张真值表"""
    if value is None:
        return None

    if isinstance(value, str):
        # 用户从别处粘贴过来的值经常带前后空格和不换行空格（U+00A0）。
        # 不 strip 的话「张三 」和「张三」是两个不同的用户名，而肉眼一样
        return value.replace('\u00a0', ' ').strip() or None

    if isinstance(value, bool):
        # bool 必须在 int 之前判 —— Python 里 `isinstance(True, int)` 是 True，
        # 顺序反了 `True` 会变成字符串 '1'
        return 'true' if value else 'false'

    if isinstance(value, (int, float, Decimal)):
        return _normalize_number(value)

    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()

    if isinstance(value, dt.timedelta):
        # Excel 的「时长」格式（[h]:mm）会读成 timedelta
        return str(value)

    # 走到这里说明遇到了没预料的类型。**不要 `str(value)` 兜底** ——
    # 那正是「解析成功但值是错的」的来源
    raise _CellError(t('error.table.cell_unsupported', type=type(value).__name__))


def _normalize_number(value: float | Decimal) -> str:
    """
    数字归一。

    🔴 **超出安全范围的数字一律拒掉，不做「尽力而为」的转换。** 19 位的数字
    （比如有人把雪花 ID 贴进 Excel）在文件里已经塌成 `1.234567890123457e+18`
    了 —— 这是 Excel 自己的精度限制（所有数值都是 IEEE double），换任何解析库
    都一样。这时候 `int()` 会给出 `1234567890123456768`：一个**看起来完全正常、
    但和用户填的不是同一个数**的整数。同硬纪律 6，两边是同一个物理限制。

    ⚠️ **`is_integer()` 那一支是防库不是防 Excel，别照着它去写别的判断。**
    实测 openpyxl 读回来时会把整数值的 float 收成 int（写进去 `13800138000.0`、
    `1.0`，读出来是 `int` 的 `13800138000`、`1`），所以走这一支的只有真带小数的值。
    真正会吐出 `13800138000.0` 的是 `python-calamine`（实测同一份文件它给 float）——
    那才是「用户直接打手机号 → 存进库里多个 `.0` 尾巴」的来源。
    留着这一支是因为它一行就能挡住换库带来的这类静默偏移，
    对应的测试直接单测这张真值表，不经过 openpyxl（否则测不到）。
    """
    if isinstance(value, Decimal):
        return str(value)

    if abs(value) > _NUMBER_SAFE_MAX:
        raise _CellError(t('error.table.cell_number_unsafe'))

    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)

    return str(value)


def build_template(
    columns: list[TableColumn],
    *,
    notes: dict[str, str] | None = None,
    sheet_title: str = 'data',
    notes_title: str = 'README',
) -> bytes:
    """
    生成一份只有表头的导入模板。

    **说明文字放在第二个工作表，不放在第一个表的第二行。** 「表头 + 中文说明行 +
    数据从第 3 行开始」这种排版看着友好，但它把「数据从第几行开始」变成了一个
    约定 —— 用户删掉说明行（很多人会删）之后整份表就错位一行，而解析器无从分辨。
    第一个表里只有表头和数据，说明单独一页，这个约定就不存在了。

    :param columns: 列声明，顺序即表头顺序
    :param notes: 列 key → 说明文字（由调用方过 i18n），缺的列不写说明
    :param sheet_title: 数据表名
    :param notes_title: 说明表名
    :return: xlsx 字节
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append([column.key for column in columns])

    readme = wb.create_sheet(notes_title)
    readme.append(['column', 'required', 'note'])
    for column in columns:
        readme.append([column.key, 'Y' if column.required else '', (notes or {}).get(column.key, '')])

    # 保存时把活动页签摆回数据表 —— 否则用户打开模板看到的是说明页，
    # 而且他若直接在这个状态下保存，`wb.active` 就指向说明页了（见 _pick_sheet）
    wb.active = 0

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
